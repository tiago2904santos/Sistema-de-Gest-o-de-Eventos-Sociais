"""DOCX da versão editada de um documento (m057), a partir do HTML dela.

O DOCX normal sai do modelo `.docx` pelo docxtpl; o documento editado à mão
não tem modelo — o que vale é o HTML que o editor completo gravou. Aqui ele
vira Word com o python-docx: parágrafos com alinhamento, negrito, itálico,
sublinhado, tachado, índices, títulos, listas, tabelas e as imagens do
timbre; o cabeçalho e o rodapé editados vão para o cabeçalho e o rodapé da
seção. O desenho fino do PDF (larguras de coluna, entrelinhas medidas, CSS
institucional por classe) não tem equivalente direto no Word e fica de fora:
o DOCX editado tem o conteúdo e a formatação de texto, não a geometria.
"""

from __future__ import annotations

import re
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

from django.conf import settings

_BLOCOS = {"p", "div", "h1", "h2", "h3", "h4", "li", "blockquote"}
_ALINHAMENTO = re.compile(r"text-align:\s*(left|right|center|justify)")
# Classes do CSS institucional que dizem o alinhamento ou a ênfase do trecho.
_CLASSES_CENTRO = ("doc-titulo", "doc-cabecalho__", "doc-rodape__", "doc-assinatura", "doc-termo__assinatura")
_IMAGENS = {"brasao": "img/brasao-pcpr.png", "marca": "img/marca-pcpr.png"}


class _Montador(HTMLParser):
    """HTML (já sanitizado) → blocos: parágrafos com trechos, e tabelas."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz: list = []
        self.containers = [self.raiz]
        self.paragrafo = None
        self.formato: list[str] = []
        self.listas: list[str] = []
        self.tabelas: list[dict] = []
        self.alinhamentos: list[str | None] = []

    # Parágrafos
    def _novo_paragrafo(self, estilo=None, alinhamento=None):
        self.paragrafo = {"tipo": "p", "trechos": [], "estilo": estilo, "alinhamento": alinhamento}
        self.containers[-1].append(self.paragrafo)
        return self.paragrafo

    def _paragrafo_atual(self):
        if self.paragrafo is None:
            herdado = next((a for a in reversed(self.alinhamentos) if a), None)
            self._novo_paragrafo(alinhamento=herdado)
        return self.paragrafo

    @staticmethod
    def _alinhamento(atributos):
        dados = dict(atributos)
        achado = _ALINHAMENTO.search(dados.get("style") or "")
        if achado:
            return achado.group(1)
        classes = dados.get("class") or ""
        if any(marca in classes for marca in _CLASSES_CENTRO):
            return "center"
        return None

    def handle_starttag(self, tag, atributos):
        if tag in ("strong", "b"):
            self.formato.append("negrito")
        elif tag in ("em", "i"):
            self.formato.append("italico")
        elif tag == "u":
            self.formato.append("sublinhado")
        elif tag in ("s", "strike"):
            self.formato.append("tachado")
        elif tag == "sup":
            self.formato.append("sobrescrito")
        elif tag == "sub":
            self.formato.append("subscrito")
        elif tag == "span":
            estilo = dict(atributos).get("style") or ""
            self.formato.append("negrito" if "font-weight: bold" in estilo or "font-weight: 700" in estilo else "")
        elif tag == "br":
            self._paragrafo_atual()["trechos"].append(("\n", ()))
        elif tag == "img":
            nome = dict(atributos).get("data-imagem")
            if nome:
                self._paragrafo_atual()["trechos"].append(("img:" + nome, ()))
        elif tag in ("ul", "ol"):
            self.listas.append(tag)
            self.paragrafo = None
        elif tag in _BLOCOS:
            alinhamento = self._alinhamento(atributos)
            self.alinhamentos.append(alinhamento)
            estilo = None
            if tag in ("h1", "h2", "h3", "h4"):
                estilo = tag
            elif tag == "li":
                estilo = "List Number" if self.listas and self.listas[-1] == "ol" else "List Bullet"
            herdado = alinhamento or next((a for a in reversed(self.alinhamentos) if a), None)
            if tag == "div":
                self.paragrafo = None  # a div só separa; o texto abre o parágrafo
            else:
                self._novo_paragrafo(estilo, herdado)
        elif tag == "table":
            tabela = {"tipo": "tabela", "linhas": []}
            self.containers[-1].append(tabela)
            self.tabelas.append(tabela)
            self.paragrafo = None
        elif tag == "tr" and self.tabelas:
            self.tabelas[-1]["linhas"].append([])
        elif tag in ("td", "th") and self.tabelas:
            if not self.tabelas[-1]["linhas"]:
                self.tabelas[-1]["linhas"].append([])
            celula: list = []
            self.tabelas[-1]["linhas"][-1].append({"blocos": celula, "cabeca": tag == "th"})
            self.containers.append(celula)
            self.paragrafo = None
            if tag == "th":
                self.formato.append("negrito")
        elif tag == "hr":
            self.paragrafo = None
            self._novo_paragrafo()["trechos"].append(("—" * 20, ()))
            self.paragrafo = None

    def handle_endtag(self, tag):
        if tag in ("strong", "b", "em", "i", "u", "s", "strike", "sup", "sub", "span"):
            if self.formato:
                self.formato.pop()
        elif tag in ("ul", "ol"):
            if self.listas:
                self.listas.pop()
            self.paragrafo = None
        elif tag in _BLOCOS:
            if self.alinhamentos:
                self.alinhamentos.pop()
            self.paragrafo = None
        elif tag in ("td", "th") and len(self.containers) > 1:
            self.containers.pop()
            self.paragrafo = None
            if tag == "th" and self.formato:
                self.formato.pop()
        elif tag == "table" and self.tabelas:
            self.tabelas.pop()
            self.paragrafo = None

    def handle_data(self, dados):
        texto = re.sub(r"\s+", " ", dados)
        if not texto.strip() and self.paragrafo is None:
            return
        paragrafo = self._paragrafo_atual()
        if not paragrafo["trechos"]:
            texto = texto.lstrip()
        if texto:
            paragrafo["trechos"].append((texto, tuple(f for f in self.formato if f)))


def _blocos(html: str) -> list:
    montador = _Montador()
    montador.feed(html or "")
    montador.close()
    return montador.raiz


def _caminho_imagem(nome: str) -> Path | None:
    relativo = _IMAGENS.get(nome)
    if not relativo:
        return None
    caminho = Path(settings.BASE_DIR) / "static" / relativo
    return caminho if caminho.is_file() else None


def _escrever_paragrafo(destino, bloco):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    estilo = bloco.get("estilo")
    try:
        paragrafo = destino.add_paragraph(style=estilo if estilo and estilo.startswith("List") else None)
    except KeyError:
        paragrafo = destino.add_paragraph()
    alinhamento = {
        "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
    }.get(bloco.get("alinhamento") or "")
    if alinhamento is not None:
        paragrafo.alignment = alinhamento
    titulo = estilo in ("h1", "h2", "h3", "h4")
    for texto, formato in bloco["trechos"]:
        if texto == "\n":
            paragrafo.add_run().add_break()
            continue
        if texto.startswith("img:"):
            caminho = _caminho_imagem(texto[4:])
            if caminho is not None:
                paragrafo.add_run().add_picture(str(caminho), height=Cm(2.2) if texto == "img:brasao" else Cm(0.8))
            continue
        trecho = paragrafo.add_run(texto)
        trecho.bold = True if titulo or "negrito" in formato else None
        trecho.italic = True if "italico" in formato else None
        trecho.underline = True if "sublinhado" in formato else None
        if "tachado" in formato:
            trecho.font.strike = True
        if "sobrescrito" in formato:
            trecho.font.superscript = True
        if "subscrito" in formato:
            trecho.font.subscript = True
        if titulo:
            trecho.font.size = Pt({"h1": 14, "h2": 13, "h3": 12, "h4": 11}[estilo])
    return paragrafo


def _escrever(destino, blocos):
    for bloco in blocos:
        if bloco["tipo"] == "p":
            if bloco["trechos"]:
                _escrever_paragrafo(destino, bloco)
            continue
        linhas = [linha for linha in bloco["linhas"] if linha]
        if not linhas:
            continue
        colunas = max(len(linha) for linha in linhas)
        try:
            tabela = destino.add_table(rows=len(linhas), cols=colunas)
        except TypeError:  # cabeçalho e rodapé pedem a largura
            from docx.shared import Cm

            tabela = destino.add_table(len(linhas), colunas, Cm(16))
        try:
            tabela.style = "Table Grid"
        except (KeyError, ValueError):
            pass
        for i, linha in enumerate(linhas):
            for j, celula in enumerate(linha):
                alvo = tabela.cell(i, j)
                inicial = alvo.paragraphs[0]
                _escrever(alvo, celula["blocos"])
                # A célula nasce com um parágrafo vazio: sai se veio conteúdo.
                if len(alvo.paragraphs) > 1 and not inicial.text:
                    inicial._element.getparent().remove(inicial._element)


def regioes_para_docx(regioes: dict) -> bytes:
    """O DOCX de uma versão editada: A4, margens do modelo, cabeçalho e rodapé
    da seção com as regiões editadas, e o corpo."""
    from docx import Document
    from docx.shared import Cm, Pt

    documento = Document()
    estilo = documento.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    secao = documento.sections[0]
    secao.page_height, secao.page_width = Cm(29.7), Cm(21.0)
    secao.left_margin = secao.right_margin = Cm(2.5)
    secao.top_margin, secao.bottom_margin = Cm(2.0), Cm(2.0)
    # O documento nasce com um parágrafo vazio no corpo.
    for paragrafo in list(documento.paragraphs):
        paragrafo._element.getparent().remove(paragrafo._element)
    for nome, destino in (("cabecalho", secao.header), ("rodape", secao.footer)):
        if regioes.get(nome):
            vazio = destino.paragraphs[0] if destino.paragraphs else None
            _escrever(destino, _blocos(regioes[nome]))
            if vazio is not None and not vazio.text and len(destino.paragraphs) > 1:
                vazio._element.getparent().remove(vazio._element)
    _escrever(documento, _blocos(regioes.get("corpo", "")))
    if not documento.paragraphs and not documento.tables:
        documento.add_paragraph("")
    saida = BytesIO()
    documento.save(saida)
    return saida.getvalue()
