"""Edição completa de documento e textos dos modelos (m057).

O documento inteiro (cabeçalho, corpo, rodapé) editado à mão vira uma versão
editada: sanitizada, versionada, usada no PDF/DOCX até voltar ao modelo, com
aviso quando os dados mudam depois. Os textos-base do modelo de cada tipo
se editam na tela de modelos e valem para os próximos documentos.
"""
import json
from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from documentos.models import DocumentoArtefato, DocumentoAssinaturaVersao, DocumentoVersaoEditada, ModeloTextoDocumento
from documentos.services import edicao_completa as edicao
from documentos.services.document_blocks import conteudo_documental
from documentos.services.document_context import contexto_do_oficio
from documentos.services.pdf_renderer import renderizar_html
from documentos.services.types import DocumentoTipo
from viagens_oficios.tests.fixtures import CenarioOficioMixin


class SanitizacaoTests(SimpleTestCase):
    def test_lista_branca_tira_script_eventos_estilos_perigosos_e_enderecos(self):
        sujo = (
            '<p onclick="roubar()" style="color:red; text-align: center; background:url(javascript:x)" class="doc-x doc-editavel">'
            'Olá <a href="javascript:alert(1)">link</a><script>alert(1)</script><style>p{}</style>'
            '<iframe src="//x"></iframe><input value="1">fim<img src="http://fora/x.png" onerror="x">'
            '<img src="/static/img/brasao-pcpr.png?v=2"></p><!-- comentário -->'
        )
        limpo = edicao.sanitizar(sujo)
        self.assertEqual(limpo, '<p style="text-align: center" class="doc-x">Olá linkfim<img data-imagem="brasao"></p>')

    def test_mantem_formatacao_listas_e_tabelas(self):
        html = ('<h1>T</h1><p style="text-align: justify">a <b>b</b> <i>c</i> <u>d</u> <s>e</s><br>f</p>'
                '<ul><li>1</li></ul><ol><li>2</li></ol>'
                '<table class="doc-tabela-livre"><tr><th colspan="2">x</th></tr><tr><td rowspan="999">y</td></tr></table>')
        limpo = edicao.sanitizar(html)
        for trecho in ('<h1>T</h1>', '<b>b</b>', '<i>c</i>', '<u>d</u>', '<s>e</s>', '<br>', '<ul><li>1</li></ul>',
                       '<th colspan="2">x</th>', '<td>y</td>', 'class="doc-tabela-livre"'):
            self.assertIn(trecho, limpo)

    def test_tags_abertas_sao_fechadas_e_align_vira_estilo(self):
        self.assertEqual(edicao.sanitizar('<div align="right"><p>x'), '<div style="text-align: right"><p>x</p></div>')
        self.assertEqual(edicao.sanitizar('<p>a<script>b'), '<p>a</p>')

    def test_regioes_desconhecidas_ou_grandes_nao_passam(self):
        self.assertEqual(edicao.sanitizar_regioes({'corpo': '<p>x</p>', 'outra': '<p>y</p>'}), {'corpo': '<p>x</p>'})
        with self.assertRaises(ValueError):
            edicao.sanitizar_regioes({'corpo': 'x' * (edicao.TAMANHO_MAXIMO_REGIAO + 1)})
        with self.assertRaises(ValueError):
            edicao.sanitizar_regioes({'corpo': ['x']})

    def test_aplicar_troca_so_a_regiao_e_resolve_as_imagens(self):
        folha = '<header><!--ed:cabecalho--><img src="a"><p>X</p><!--/ed:cabecalho--></header><main><!--ed:corpo--><p>velho \\1</p><!--/ed:corpo--></main>'
        nova = edicao.aplicar_regioes(folha, {'corpo': '<p>novo \\1</p><img data-imagem="marca">'}, {'marca': '/static/m.png'})
        self.assertIn('<p>X</p>', nova)
        self.assertIn('<!--ed:corpo--><p>novo \\1</p><img src="/static/m.png" data-imagem="marca"><!--/ed:corpo-->', nova)
        self.assertEqual(edicao.extrair_regioes(nova)['corpo'], '<p>novo \\1</p><img src="/static/m.png" data-imagem="marca">')
        # O hash olha o texto, não o CSS nem as marcas.
        self.assertEqual(edicao.impressao_digital(folha), edicao.impressao_digital(folha.replace('<p>', '<p class="z">')))
        self.assertNotEqual(edicao.impressao_digital(folha), edicao.impressao_digital(nova))

    def test_docx_da_versao_editada(self):
        from docx import Document

        from documentos.services.html_docx import regioes_para_docx

        conteudo = regioes_para_docx({
            'cabecalho': '<img data-imagem="brasao"><p>POLÍCIA CIVIL</p>',
            'corpo': '<h1>Título</h1><p style="text-align: justify">Texto <b>forte</b></p><ul><li>item</li></ul>'
                     '<table><tr><td>a</td><td>b</td></tr></table>',
            'rodape': '<p>Rodapé livre</p>',
        })
        documento = Document(BytesIO(conteudo))
        textos = [p.text for p in documento.paragraphs]
        self.assertEqual(textos[:3], ['Título', 'Texto forte', 'item'])
        self.assertEqual([c.text for c in documento.tables[0].rows[0].cells], ['a', 'b'])
        self.assertIn('POLÍCIA CIVIL', [p.text for p in documento.sections[0].header.paragraphs])
        self.assertIn('Rodapé livre', [p.text for p in documento.sections[0].footer.paragraphs])


class EdicaoCompletaTests(CenarioOficioMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.oficio = self.criar()

    def url(self, nome, chave='oficio', pk=None, v='', *args):
        url = reverse(f'documentos:{nome}', args=[chave, pk or self.oficio.pk, *args])
        return f'{url}?v={v}' if v else url

    def salvar(self, regioes, *, chave='oficio', pk=None, v='', estado=None):
        if estado is None:
            atual = DocumentoVersaoEditada.objects.order_by('-criado_em', '-pk').filter(tipo_documento=DocumentoTipo(
                {'termo_oficio': 'termo_autorizacao'}.get(chave, chave)).value, variante=v).first()
            estado = atual.pk if atual else ''
        return self.client.post(self.url('editor_completo_salvar', chave, pk, v), data=json.dumps({'estado': estado, 'regioes': regioes}),
                                content_type='application/json')

    def html_pdf(self):
        return renderizar_html(DocumentoTipo.OFICIO, contexto_do_oficio(self.oficio, modo='pdf'), modo='pdf')

    def test_pagina_e_folha_abrem_com_as_regioes_marcadas(self):
        r = self.client.get(self.url('editor_completo'))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'data-dcp-editavel')
        folha = self.client.get(self.url('editor_completo_folha')).content.decode()
        for regiao in ('cabecalho', 'corpo', 'rodape'):
            self.assertIn(f'<!--ed:{regiao}-->', folha)
        self.assertNotIn('data-doc-campo', folha)  # a folha do editor completo não é a de campos
        self.assertIn('Missão F4', folha)
        # O editor de campos oferece o editor completo.
        self.assertContains(self.client.get(self.url('editor_embutido')), self.url('editor_completo'))

    def test_salvar_grava_sanitizado_e_o_pdf_sai_da_versao_editada(self):
        r = self.salvar({'corpo': '<p>Texto <b>livre</b> do ofício<script>x()</script></p>', 'cabecalho': '<p onclick="x">CABEÇALHO NOVO</p>'})
        self.assertEqual(r.status_code, 200, r.content)
        versao = DocumentoVersaoEditada.objects.get()
        self.assertEqual(versao.regioes['corpo'], '<p>Texto <b>livre</b> do ofício</p>')
        self.assertEqual(versao.regioes['cabecalho'], '<p>CABEÇALHO NOVO</p>')
        self.assertEqual(versao.criado_por, self.user)
        self.assertTrue(versao.impressao_base)
        html = self.html_pdf()
        self.assertIn('Texto <b>livre</b> do ofício', html)
        self.assertIn('CABEÇALHO NOVO', html)
        self.assertNotIn('Missão F4', html)
        # O rodapé não foi mandado: continua o do modelo.
        self.assertIn('doc-rodape__unidade', html)
        # Entra no payload de geração (e, por ele, na chave de cache).
        self.assertEqual(conteudo_documental(DocumentoTipo.OFICIO, self.oficio)['versao_editada']['id'], versao.pk)
        with self.assertRaises(ValidationError):
            versao.save()  # imutável

    def test_geracao_real_usa_a_versao_editada_no_pdf_e_no_docx(self):
        from docx import Document

        from documentos.services.types import DocumentoFormato
        from viagens_oficios.document_generation import gerar_documento

        self.salvar({'corpo': '<p>Corpo editado para o DOCX</p>'})
        docx = gerar_documento(self.oficio, DocumentoFormato.DOCX)
        self.assertIn('Corpo editado para o DOCX', [p.text for p in Document(BytesIO(docx.conteudo)).paragraphs])
        artefato = DocumentoArtefato.objects.get(pk=docx.artefato_id)
        self.assertIn('versao_editada', artefato.payload_snapshot['documento'])

    def test_conflito_quando_outra_pessoa_gravou_antes(self):
        self.assertEqual(self.salvar({'corpo': '<p>um</p>'}).status_code, 200)
        r = self.salvar({'corpo': '<p>dois</p>'}, estado='')
        self.assertEqual(r.status_code, 409)
        self.assertEqual(DocumentoVersaoEditada.objects.count(), 1)

    def test_aviso_de_desatualizada_quando_os_dados_mudam(self):
        from documentos.editor.completo import situacao_da_edicao
        from documentos.editor.vinculos import vinculo_do_tipo

        self.salvar({'corpo': '<p>editado</p>'})
        vinculo = vinculo_do_tipo('oficio')
        self.assertFalse(situacao_da_edicao(vinculo, vinculo.carregar(self.oficio.pk))['desatualizada'])
        type(self.oficio).objects.filter(pk=self.oficio.pk).update(motivo='Outro motivo depois da edição')
        self.assertTrue(situacao_da_edicao(vinculo, vinculo.carregar(self.oficio.pk))['desatualizada'])
        self.assertContains(self.client.get(self.url('editor_completo')), 'Os dados mudaram depois desta edição')
        self.assertContains(self.client.get(self.url('editor_embutido')), 'Os dados mudaram depois da edição completa')

    def test_voltar_ao_modelo_e_restaurar_guardam_o_historico(self):
        self.salvar({'corpo': '<p>primeira</p>'})
        primeira = DocumentoVersaoEditada.objects.get()
        self.salvar({'corpo': '<p>segunda</p>'})
        r = self.client.post(self.url('editor_completo_modelo'))
        self.assertEqual(r.status_code, 302)
        self.assertIsNone(edicao.vigente(DocumentoTipo.OFICIO, self.oficio))
        self.assertIn('Missão F4', self.html_pdf())
        r = self.client.post(self.url('editor_completo_restaurar', 'oficio', None, '', primeira.pk))
        self.assertEqual(r.status_code, 302)
        vigente = edicao.vigente(DocumentoTipo.OFICIO, self.oficio)
        self.assertEqual((vigente.acao, vigente.restaurada_de, vigente.regioes['corpo']), ('restauracao', primeira, '<p>primeira</p>'))
        self.assertIn('primeira', self.html_pdf())
        self.assertEqual(DocumentoVersaoEditada.objects.count(), 4)
        pagina = self.client.get(self.url('editor_completo'))
        self.assertContains(pagina, 'Voltou ao modelo')
        self.assertContains(pagina, 'Restaurar esta versão')

    def test_quem_so_consulta_ve_mas_nao_grava(self):
        leitor = get_user_model().objects.create_user(username='leitor-m057', deve_trocar_senha=False)
        leitor.setores.add(self.setor)
        self.client.force_login(leitor)
        pagina = self.client.get(self.url('editor_completo'))
        self.assertEqual(pagina.status_code, 200)
        self.assertNotContains(pagina, 'data-dcp-editavel')
        self.assertEqual(self.salvar({'corpo': '<p>x</p>'}).status_code, 403)
        self.assertEqual(self.client.post(self.url('editor_completo_modelo')).status_code, 403)
        self.assertFalse(DocumentoVersaoEditada.objects.exists())

    def test_documento_assinado_so_se_edita_reabrindo(self):
        artefato = DocumentoArtefato.objects.create(tipo='oficio', formato='pdf', oficio=self.oficio, hash_sha256='0' * 64,
                                                    arquivo=ContentFile(b'%PDF-1.4', name='o.pdf'))
        versao = DocumentoAssinaturaVersao.objects.create(artefato=artefato, arquivo=ContentFile(b'%PDF-1.4 a', name='a.pdf'),
                                                          hash_sha256='1' * 64)
        r = self.salvar({'corpo': '<p>x</p>'})
        self.assertEqual(r.status_code, 403)
        self.assertIn('versão assinada', r.json()['mensagem'])
        self.assertContains(self.client.get(self.url('editor_completo')), 'Só leitura')
        versao.revogada_em = versao.criado_em
        versao.save()
        self.assertEqual(self.salvar({'corpo': '<p>x</p>'}).status_code, 200)

    def test_cancelado_nao_edita(self):
        type(self.oficio).objects.filter(pk=self.oficio.pk).update(cancelado=True)
        self.assertEqual(self.salvar({'corpo': '<p>x</p>'}).status_code, 403)

    def test_termo_de_um_servidor_nao_muda_o_do_outro(self):
        from viagens_termos.models import TermoAutorizacao

        termo = TermoAutorizacao.objects.create(oficio=self.oficio, destino_estado=self.uf, destino_cidade=self.destino,
                                                data_evento_inicio=date(2026, 9, 10), viatura=self.viatura)
        termo.servidores.set([self.a, self.b])
        r = self.salvar({'corpo': '<p>Termo só da Ana</p>'}, chave='termo_autorizacao', pk=termo.pk, v=str(self.a.pk))
        self.assertEqual(r.status_code, 200, r.content)
        folha_a = self.client.get(self.url('editor_completo_folha', 'termo_autorizacao', termo.pk, str(self.a.pk))).content.decode()
        folha_b = self.client.get(self.url('editor_completo_folha', 'termo_autorizacao', termo.pk, str(self.b.pk))).content.decode()
        self.assertIn('Termo só da Ana', folha_a)
        self.assertNotIn('Termo só da Ana', folha_b)
        # A geração do termo do cadastro usa a mesma variante do editor.
        from viagens_termos.services import build_termo_cadastro_payload, _conteudo_documental, variante_do_termo_cadastro

        self.assertIn('versao_editada', _conteudo_documental(termo, variante_do_termo_cadastro(self.a)))
        self.assertNotIn('versao_editada', _conteudo_documental(termo, variante_do_termo_cadastro(self.b)))
        self.assertTrue(build_termo_cadastro_payload(termo, self.a))

    def test_todos_os_documentos_de_viagens_abrem_no_editor_completo(self):
        from documentos.editor.vinculos import VINCULOS

        for chave in ('oficio', 'justificativa', 'termo_oficio'):
            with self.subTest(chave=chave):
                v = str(self.a.pk) if chave == 'termo_oficio' else ''
                folha = self.client.get(self.url('editor_completo_folha', chave, self.oficio.pk, v))
                self.assertEqual(folha.status_code, 200)
                self.assertIn('<!--ed:corpo-->', folha.content.decode())
        # Todo vínculo registrado tem a folha completa (os do Coffee Break incluídos).
        for vinculo in VINCULOS.values():
            self.assertTrue(callable(getattr(vinculo, 'html_documento', None)))


class ModelosDeTextoTests(CenarioOficioMixin, TestCase):
    """A tela de modelos: a lista dos tipos e, de cada um, o documento montado
    com os textos do modelo editáveis no lugar (sem caixas de texto)."""

    def setUp(self):
        super().setUp()
        self.oficio = self.criar()
        self.user.groups.add(Group.objects.get(name='VIAGENS_GESTOR'))
        self.url = reverse('documentos:modelos_tipo', args=['termo_autorizacao'])

    def html_termo(self):
        from documentos.services.document_context import contexto_de_payload
        from viagens_oficios.documents import build_termo_payload
        from viagens_termos.services import _conteudo_documental, _legacy_docx_context

        payload = build_termo_payload(self.oficio, self.a)
        payload['documento'] = _conteudo_documental(self.oficio, str(self.a.pk))
        contexto = contexto_de_payload(DocumentoTipo.TERMO_AUTORIZACAO, payload, _legacy_docx_context(payload))
        return renderizar_html(DocumentoTipo.TERMO_AUTORIZACAO, contexto, modo='pdf')

    def salvar(self, tipo, blocos, estado=None):
        from documentos.services import modelos_texto

        if estado is None:
            estado = modelos_texto.estado(tipo)
        return self.client.post(reverse('documentos:modelos_salvar', args=[tipo]),
                                data=json.dumps({'estado': estado, 'blocos': blocos}), content_type='application/json')

    def folha(self, tipo):
        r = self.client.get(reverse('documentos:modelos_folha', args=[tipo]))
        self.assertEqual(r.status_code, 200)
        return r.content.decode()

    def test_sem_alteracao_o_termo_sai_com_o_texto_de_sempre(self):
        html = self.html_termo()
        self.assertIn('manifesto o interesse em participar do PCPR na Comunidade, <strong>', html)
        self.assertIn('</strong> para execução de atividades inerentes à Assessoria de Comunicação Social - ASCOM/PCPR.', html)

    def test_lista_dos_tipos_com_abrir_modelo(self):
        r = self.client.get(reverse('documentos:modelos_modulo', args=['viagens']))
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, '<textarea')
        for rotulo in ('Ofício', 'Justificativa', 'Termo de autorização', 'Ordem de serviço', 'Plano de trabalho',
                       'Relatório técnico', 'Diário de bordo'):
            self.assertContains(r, rotulo)
        self.assertContains(r, 'Abrir modelo', count=7)
        self.assertContains(r, reverse('documentos:modelos_tipo', args=['termo_autorizacao']))
        self.assertContains(r, 'Nunca alterado', count=7)
        self.assertEqual(self.salvar('termo_autorizacao', {'titulo': 'Termo de autorização F4'}).status_code, 200)
        r = self.client.get(reverse('documentos:modelos'))
        self.assertContains(r, '<b>1 de 6</b> texto personalizado')
        self.assertContains(r, 'Nunca alterado', count=6)
        self.assertContains(r, 'operador-f4')

    def test_oficio_abre_montado_do_documento_real_com_os_blocos_editaveis(self):
        r = self.client.get(reverse('documentos:modelos_tipo', args=['oficio']))
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, '<textarea')
        self.assertContains(r, 'ds-v32-bridge.css')
        self.assertContains(r, f'a partir de <b>Ofício {self.oficio.numero_formatado}')
        self.assertContains(r, 'Salvar modelo')
        self.assertContains(r, 'próximos documentos')
        folha = self.folha('oficio')
        # Cada bloco do modelo é um trecho editável; o resto (dados, tabela da equipe) sai como no documento.
        for chave in ('secretaria', 'abertura', 'fecho', 'declaracao_cartao'):
            self.assertIn(f'data-mod-bloco="{chave}"', folha)
        self.assertEqual(folha.count('contenteditable="true"'), 4)
        self.assertIn('<span class="mod-chip" contenteditable="false" data-mod-campo="assunto"', folha)
        self.assertIn('>Assunto</span>', folha)
        self.assertIn('Ana Teste', folha)
        self.assertIn('brasao', folha)
        self.assertNotIn('', folha)
        self.assertNotIn('data-doc-campo', folha)

    def test_termo_sem_documento_abre_com_exemplo_e_campos_em_negrito(self):
        r = self.client.get(self.url)
        self.assertContains(r, 'dados fictícios')
        folha = self.folha('termo_autorizacao')
        self.assertIn('SERVIDOR DE EXEMPLO', folha)
        self.assertIn('data-mod-bloco="texto"', folha)
        self.assertIn('class="mod-chip mod-chip--negrito" contenteditable="false" data-mod-campo="periodo"', folha)
        self.assertIn('data-mod-campo="destino"', folha)
        self.assertIn('>Período</span>', folha)

    def test_salvar_grava_so_o_que_mudou_e_o_termo_sai_com_o_texto_novo(self):
        from documentos.editor.blocos import bloco

        novo = 'declaro participar da Operação Verão em {destino}, {periodo}, lotado em {unidade}.'
        r = self.salvar('termo_autorizacao', {
            'texto': novo + '\n',  # o que a folha manda: quebras e espaços sobrando saem
            'titulo': bloco(DocumentoTipo.TERMO_AUTORIZACAO, 'titulo').padrao,  # igual ao de hoje: não grava
        })
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()['gravados'], ['texto'])
        linha = ModeloTextoDocumento.objects.get()
        self.assertEqual((linha.tipo_documento, linha.chave, linha.texto, linha.criado_por),
                         ('termo_autorizacao', 'texto', novo, self.user))
        html = self.html_termo()
        self.assertIn('declaro participar da Operação Verão em <strong>Londrina/PR</strong>, <strong>', html)
        self.assertIn('lotado em UNIDADE F4.', html)
        self.assertNotIn('PCPR na Comunidade', html)
        # A folha mostra o texto novo com os campos como etiquetas e o bloco marcado como personalizado.
        folha = self.folha('termo_autorizacao')
        self.assertIn('data-mod-alterado="1"', folha)
        self.assertIn('Operação Verão em <span class="mod-chip mod-chip--negrito"', folha)
        # Mandar o mesmo texto de novo não grava nada.
        self.assertEqual(self.salvar('termo_autorizacao', {'texto': novo}).json()['gravados'], [])
        self.assertEqual(ModeloTextoDocumento.objects.count(), 1)

    def test_campo_que_o_bloco_nao_aceita_e_recusado(self):
        r = self.salvar('termo_autorizacao', {'texto': 'em {destino} e {xyz}', 'titulo': 'Termo novo'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('{xyz}', r.json()['mensagem'])
        self.assertIn('Destino', r.json()['mensagem'])
        r = self.salvar('termo_autorizacao', {'titulo': 'Termo de {destino}'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('aceita: nenhum', r.json()['mensagem'])
        self.assertEqual(self.salvar('termo_autorizacao', {'nao_existe': 'x'}).status_code, 400)
        self.assertEqual(self.salvar('termo_autorizacao', {'titulo': ['x']}).status_code, 400)
        self.assertFalse(ModeloTextoDocumento.objects.exists())

    def test_voltar_ao_padrao_e_usar_um_texto_do_historico(self):
        from documentos.services import modelos_texto

        self.salvar('termo_autorizacao', {'texto': 'declaro participar da Operação Verão em {destino}.'})
        linha = ModeloTextoDocumento.objects.get()
        r = self.client.get(self.url)
        self.assertContains(r, 'Histórico')
        self.assertContains(r, 'Texto da manifestação')
        r = self.client.post(self.url, {'padrao': 'texto', 'estado': modelos_texto.estado('termo_autorizacao')})
        self.assertEqual(r.status_code, 302)
        self.assertIn('PCPR na Comunidade', self.html_termo())
        self.assertEqual(ModeloTextoDocumento.objects.count(), 2)
        self.assertNotIn('data-mod-alterado', self.folha('termo_autorizacao'))
        self.assertContains(self.client.get(self.url), 'Usar este texto de novo')
        self.client.post(self.url, {'usar': linha.pk, 'estado': modelos_texto.estado('termo_autorizacao')})
        self.assertIn('Operação Verão', self.html_termo())

    def test_duas_pessoas_salvando_ao_mesmo_tempo(self):
        from documentos.services import modelos_texto

        lida = modelos_texto.estado('termo_autorizacao')
        self.assertEqual(self.salvar('termo_autorizacao', {'titulo': 'Primeira pessoa'}, estado=lida).status_code, 200)
        r = self.salvar('termo_autorizacao', {'titulo': 'Segunda pessoa'}, estado=lida)
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.json()['conflito'])
        self.assertIn('Recarregue', r.json()['mensagem'])
        self.assertEqual(modelos_texto.texto_vigente('termo_autorizacao', 'titulo'), 'Primeira pessoa')
        # Voltar ao padrão com a tela velha também não passa.
        self.client.post(self.url, {'padrao': 'titulo', 'estado': lida})
        self.assertEqual(modelos_texto.texto_vigente('termo_autorizacao', 'titulo'), 'Primeira pessoa')

    def test_texto_do_modelo_nao_passa_por_cima_do_reescrito_no_documento_nem_da_versao_editada(self):
        from documentos.services.document_blocks import gravar_override

        self.assertEqual(self.salvar('oficio', {'fecho': 'Atenciosamente,'}).status_code, 200)
        self.assertIn('Atenciosamente,', renderizar_html(DocumentoTipo.OFICIO, contexto_do_oficio(self.oficio, modo='pdf'), modo='pdf'))
        gravar_override(DocumentoTipo.OFICIO, self.oficio, 'fecho', 'Cordialmente,', self.user)
        html = renderizar_html(DocumentoTipo.OFICIO, contexto_do_oficio(self.oficio, modo='pdf'), modo='pdf')
        self.assertIn('Cordialmente,', html)
        self.assertNotIn('Atenciosamente,', html)
        # A folha do modelo mostra o texto do modelo, não o reescrito só neste ofício.
        folha = self.folha('oficio')
        self.assertIn('Atenciosamente,', folha)
        self.assertNotIn('Cordialmente,', folha)

    def test_todos_os_tipos_de_viagens_abrem_com_todos_os_blocos(self):
        from documentos.editor.blocos import REGISTRO_BLOCOS
        from documentos.editor.modelos import tipos_do_usuario

        tipos = tipos_do_usuario(self.user)
        self.assertEqual({t['valor'] for t in tipos}, {'oficio', 'justificativa', 'termo_autorizacao', 'ordem_servico',
                                                      'plano_trabalho', 'relatorio_tecnico', 'diario_bordo'})
        for t in tipos:
            with self.subTest(tipo=t['valor']):
                pagina = self.client.get(t['url'])
                self.assertEqual(pagina.status_code, 200)
                self.assertNotContains(pagina, '<textarea')
                folha = self.folha(t['valor'])
                self.assertNotIn('', folha)
                for chave in REGISTRO_BLOCOS[t['tipo']]:
                    self.assertTrue(f'data-mod-bloco="{chave}"' in folha or f'data-mod-bloco="{chave}"' in pagina.content.decode(), chave)

    def test_so_o_gestor_administra_os_modelos(self):
        self.user.groups.remove(Group.objects.get(name='VIAGENS_GESTOR'))
        self.assertEqual(self.client.get(reverse('documentos:modelos')).status_code, 403)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.client.get(reverse('documentos:modelos_folha', args=['termo_autorizacao'])).status_code, 403)
        self.assertEqual(self.client.post(self.url, {'padrao': 'texto'}).status_code, 403)
        self.assertEqual(self.salvar('termo_autorizacao', {'titulo': 'x'}).status_code, 403)
        self.assertFalse(ModeloTextoDocumento.objects.exists())
        self.assertEqual(self.client.get(reverse('documentos:modelos_tipo', args=['nao_existe'])).status_code, 404)
        # Os do Coffee Break são da administração do Coffee Break.
        self.user.groups.add(Group.objects.get(name='VIAGENS_GESTOR'))
        self.assertEqual(self.client.get(reverse('documentos:modelos_tipo', args=['coffee_break_certifico'])).status_code, 403)
        self.assertEqual(self.salvar('coffee_break_certifico', {'cb_titulo': 'x'}).status_code, 403)
