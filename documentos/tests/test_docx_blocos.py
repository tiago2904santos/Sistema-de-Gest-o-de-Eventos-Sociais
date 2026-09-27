"""O DOCX sai com os parágrafos reescritos e as quebras de página do
documento, como o PDF (m111)."""
import json
from io import BytesIO

from django.test import TestCase
from django.urls import reverse

from documentos.editor.blocos import BLOCOS_OFICIO, BLOCOS_TERMO
from documentos.services.docx_blocos import aplicar_conteudo_documental
from documentos.services.document_context import contexto_do_oficio
from documentos.services.pdf_renderer import renderizar_html
from documentos.services.types import DocumentoFormato, DocumentoTipo
from viagens_oficios.tests.fixtures import CenarioOficioMixin

PADRAO = {b.chave: b.padrao for b in BLOCOS_OFICIO}
PADRAO_TERMO = {b.chave: b.padrao for b in BLOCOS_TERMO}


def _textos(docx_bytes):
    from docx import Document

    documento = Document(BytesIO(docx_bytes))
    textos = [p.text.strip() for p in documento.paragraphs]
    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                textos.extend(p.text.strip() for p in celula.paragraphs)
    for secao in documento.sections:
        textos.extend(p.text.strip() for p in secao.header.paragraphs)
    return textos


def _quebras(docx_bytes):
    from docx import Document

    documento = Document(BytesIO(docx_bytes))
    return documento.element.body.xpath('.//w:br[@w:type="page"]')


class DocxComBlocosTests(CenarioOficioMixin, TestCase):
    def _patch_bloco(self, o, chave, conteudo):
        url = reverse('documentos:editor_bloco', args=['oficio', o.pk, chave])
        r = self.client.patch(url, data=json.dumps({'versao': '', 'valores': {'conteudo': conteudo}}), content_type='application/json')
        self.assertEqual(r.status_code, 200, r.content)

    def _quebra(self, o, chave):
        url = reverse('documentos:editor_quebra', args=['oficio', o.pk, chave])
        self.assertEqual(self.client.patch(url, data='{"ativa": true}', content_type='application/json').status_code, 200)

    def test_docx_do_oficio_sai_com_os_textos_editados_e_as_quebras(self):
        from viagens_oficios.document_generation import gerar_documento

        o = self.criar()
        self._patch_bloco(o, 'abertura', 'Senhor Delegado, venho solicitar {assunto} e o custeio,\nconforme o cronograma:')
        self._patch_bloco(o, 'fecho', 'Atenciosamente,')
        self._patch_bloco(o, 'declaracao_cartao', 'Declaro que todos têm cartão corporativo apto.')
        self._patch_bloco(o, 'secretaria', 'SECRETARIA DA SEGURANÇA')
        self._quebra(o, 'antes_assinatura')
        self._quebra(o, 'apos_equipe')
        resultado = gerar_documento(o, DocumentoFormato.DOCX)
        textos = _textos(resultado.conteudo)
        junto = '\n'.join(textos)
        # O marcador {assunto} recebe o que o modelo tinha no lugar dele.
        self.assertIn('Senhor Delegado, venho solicitar autorização e o custeio,\nconforme o cronograma:', textos)
        self.assertIn('Atenciosamente,', textos)
        self.assertIn('Declaro que todos têm cartão corporativo apto.', textos)
        self.assertIn('SECRETARIA DA SEGURANÇA', textos)
        for chave in ('fecho', 'declaracao_cartao', 'secretaria'):
            self.assertNotIn(PADRAO[chave], junto)
        self.assertNotIn('Através deste, solicito', junto)
        self.assertEqual(len(_quebras(resultado.conteudo)), 2)
        # O que o DOCX diz é o que o PDF diz.
        html = renderizar_html(DocumentoTipo.OFICIO, contexto_do_oficio(o, modo='pdf'), modo='pdf')
        self.assertIn('Atenciosamente,', html)
        self.assertIn('Declaro que todos têm cartão corporativo apto.', html)
        self.assertIn('venho solicitar autorização e o custeio,<br>conforme o cronograma:', html)

    def test_docx_sem_alteracao_sai_como_sempre(self):
        from viagens_oficios.document_generation import gerar_documento

        o = self.criar()
        resultado = gerar_documento(o, DocumentoFormato.DOCX)
        textos = _textos(resultado.conteudo)
        self.assertIn(PADRAO['fecho'], textos)
        self.assertIn(PADRAO['declaracao_cartao'], textos)
        self.assertEqual(len(_quebras(resultado.conteudo)), 0)

    def test_docx_do_termo_troca_o_texto_da_manifestacao_com_os_marcadores(self):
        from viagens_termos.services import gerar_termo_um

        o = self.criar()
        url = reverse('documentos:editor_bloco', args=['termo_oficio', o.pk, 'texto']) + f'?v={self.a.pk}'
        novo = 'manifesto interesse em atuar no PCPR na Comunidade, {periodo}, em {destino}, pela unidade {unidade}.'
        r = self.client.patch(url, data=json.dumps({'versao': '', 'valores': {'conteudo': novo}}), content_type='application/json')
        self.assertEqual(r.status_code, 200, r.content)
        resultado = gerar_termo_um(o, self.a, DocumentoFormato.DOCX)
        junto = '\n'.join(_textos(resultado.conteudo))
        self.assertIn('manifesto interesse em atuar no PCPR na Comunidade, ', junto)
        self.assertIn('na Comunidade, nos dias 10 até 11 de setembro de 2026, em Londrina/PR, pela unidade UNIDADE F4.', junto)
        self.assertNotIn('para execução de atividades inerentes', junto)

    def test_sem_conteudo_documental_os_bytes_voltam_iguais(self):
        self.assertEqual(aplicar_conteudo_documental(DocumentoTipo.OFICIO, b'abc', None), b'abc')
        self.assertEqual(aplicar_conteudo_documental('coffee_break_oficio', b'abc', {'blocos': {'x': {'conteudo': 'y'}}}), b'abc')
        self.assertEqual(aplicar_conteudo_documental(DocumentoTipo.OFICIO, b'abc', {'blocos': {'fecho': {'conteudo': PADRAO['fecho']}}, 'quebras': []}), b'abc')
