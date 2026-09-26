"""Meta 3 — paridade das telas de ofícios com o Gerenciador de Viagens.

Cobre a lista no padrão das listas de termos e justificativas (busca por
número, protocolo, motivo ou destino; situações na trilha com contagem; uma
célula por ofício; o menu único e o modal "Baixar documentos"); o formulário por blocos, com
a conferência, os documentos, o histórico e o encerramento na mesma tela; e os
catálogos no padrão dos cadastros.
"""

from datetime import datetime, timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import Modulo, Setor
from cadastros.models import Estado, Municipio, Regiao
from viagens_cadastros.models import Cargo, Servidor, Unidade, Viatura
from viagens_oficios.models import ModeloJustificativa, ModeloMotivoOficio, Oficio
from viagens_oficios.justificativas_services import get_or_create_justificativa_oficio
from viagens_oficios.services import reservar_numero_oficio
from viagens_roteiros.models import Roteiro, RoteiroDestino, RoteiroTrecho


class Cenario(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="operador-m3", deve_trocar_senha=False)
        setor = Setor.objects.create(nome="Setor M3")
        Modulo.objects.get(codigo="VIAGENS").setores.add(setor)
        self.user.setores.add(setor)
        self.user.groups.add(Group.objects.get(name="VIAGENS_OPERADOR"))
        self.client.force_login(self.user)

        regiao = Regiao.objects.get_or_create(nome="Interior")[0]
        self.pr = Estado.objects.get_or_create(sigla="PR", defaults={"nome": "Paraná", "codigo_ibge": 41})[0]
        self.curitiba = Municipio.objects.get_or_create(codigo_ibge=4106902, defaults={"nome": "Curitiba", "estado": self.pr, "regiao": regiao})[0]
        self.antonina = Municipio.objects.get_or_create(codigo_ibge=4101309, defaults={"nome": "Antonina", "estado": self.pr, "regiao": regiao})[0]
        cargo = Cargo.objects.create(nome="AGENTE DE POLÍCIA JUDICIÁRIA")
        self.ascom = Unidade.objects.create(nome="ASSESSORIA DE COMUNICAÇÃO", sigla="ASCOM")
        self.janine = Servidor.objects.create(nome="JANINE LACERDA DO PRADO", cargo=cargo, unidade=self.ascom, cpf="11111111111")
        self.joao = Servidor.objects.create(nome="JOÃO MARIO DE GOES", cargo=cargo, unidade=self.ascom, cpf="22222222222")
        self.duster = Viatura.objects.create(placa="AAA1234", modelo="DUSTER")
        self.hoje = timezone.localdate()

    def roteiro(self, dias, destino=None, valor=Decimal("4358.25")):
        saida = timezone.make_aware(datetime.combine(self.hoje + timedelta(days=dias), datetime.min.time()).replace(hour=8))
        r = Roteiro.objects.create(origem_municipio=self.curitiba, saida_dt=saida, retorno_chegada_dt=saida + timedelta(days=5, hours=1),
                                   valor_diarias=valor, resumo_diarias="5 x 100%", quantidade_servidores=2,
                                   valor_diarias_extenso="quatro mil trezentos e cinquenta e oito reais e vinte e cinco centavos")
        RoteiroDestino.objects.create(roteiro=r, municipio=destino or self.antonina, ordem=1)
        RoteiroTrecho.objects.create(roteiro=r, ordem=1, origem_municipio=self.curitiba, destino_municipio=destino or self.antonina,
                                     saida_dt=saida, chegada_dt=saida + timedelta(minutes=45))
        return r

    def oficio(self, dias=None, protocolo="", servidores=(), motorista=None, viatura=None, justificativa="", cancelar=False, data_criacao=None):
        o = Oficio(data_criacao=data_criacao or self.hoje, protocolo=protocolo, motivo="Cobertura jornalística",
                   roteiro=self.roteiro(dias) if dias is not None else None, viatura=viatura, motorista=motorista)
        o.diarias_quantidade_servidores = len(servidores) or None
        o.save()
        o.servidores.set(servidores)
        o.servidores_termo_autorizacao.set(servidores)
        reservar_numero_oficio(o, ano=o.data_criacao.year)
        j = get_or_create_justificativa_oficio(o)
        if justificativa:
            j.texto = justificativa
            j.save()
        if cancelar:
            o.cancelar("Adiado")
        return o

    def lista(self, **params):
        return self.client.get(reverse("viagens_oficios:lista"), params)


class ListaOficiosTests(Cenario):
    def test_linha_no_padrao_das_listas_de_termos(self):
        o = self.oficio(dias=-20, protocolo="123456789", servidores=[self.janine, self.joao], motorista=self.joao, viatura=self.duster, justificativa="teste 1")
        r = self.lista()
        self.assertTemplateUsed(r, "components/v32/cad_rail.html")
        self.assertContains(r, 'class="pa-card pa-card--lista cad-lista"')
        self.assertContains(r, f"Nº {o.numero_formatado} · Protocolo 12.345.678-9")
        self.assertContains(r, "ANTONINA/PR")
        self.assertContains(r, "há 15 dias")
        self.assertContains(r, "JANINE LACERDA DO PRADO, JOÃO MARIO DE GOES (motorista)")
        self.assertContains(r, "DUSTER · AAA-1234")
        self.assertContains(r, "R$ 4.358,25 · 5 x 100%")
        self.assertContains(r, "Justificativa preenchida")
        # O cartão e os filtros antigos saíram.
        for marca in ["of-cartao__rodape", "of-catalogos", "Mostrando <strong>"]:
            self.assertNotContains(r, marca)

    def test_linha_mostra_o_tipo_do_oficio(self):
        self.oficio(dias=-2, servidores=[self.janine])
        r = self.lista()
        self.assertContains(r, ">Convalidação</span>")
        self.assertContains(r, "antes da data do ofício")

    def test_filtros_prontos_e_ordenacao(self):
        perto = self.oficio(dias=3)
        longe = self.oficio(dias=40)
        antigo = self.oficio(dias=-30)
        de = (self.hoje + timedelta(days=1)).isoformat()
        ate = (self.hoje + timedelta(days=15)).isoformat()
        r = self.lista(viagem_de=de, viagem_ate=ate)
        self.assertEqual([l["oficio"].pk for l in r.context["linhas"]], [perto.pk])
        self.assertContains(r, f"Viagem a partir de: {(self.hoje + timedelta(days=1)):%d/%m/%Y}")
        self.assertContains(r, 'name="viagem_de"')
        r = self.lista(sort="viagem_asc")
        self.assertEqual([l["oficio"].pk for l in r.context["linhas"]], [antigo.pk, perto.pk, longe.pk])
        self.assertContains(r, "Ordem: Viagem: mais próxima")
        # Filtro inválido é ignorado, não derruba a página.
        r = self.lista(viagem_de="31/02", ano="abc")
        self.assertEqual(len(r.context["linhas"]), 3)
        # A exportação leva o mesmo recorte.
        self.assertContains(r, reverse("viagens_oficios:exportar"))

    def test_exportar_excel_com_o_recorte(self):
        import io
        from openpyxl import load_workbook
        perto = self.oficio(dias=3, protocolo="123456789", servidores=[self.janine])
        self.oficio(dias=40)
        de = (self.hoje + timedelta(days=1)).isoformat()
        ate = (self.hoje + timedelta(days=15)).isoformat()
        r = self.client.get(reverse("viagens_oficios:exportar"), {"viagem_de": de, "viagem_ate": ate})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        aba = load_workbook(io.BytesIO(r.content)).active
        linhas = list(aba.iter_rows(values_only=True))
        self.assertEqual(linhas[0][:5], ("Nº", "Data do ofício", "Protocolo", "Situação", "Tipo"))
        self.assertEqual(len(linhas), 2)
        self.assertEqual(linhas[1][0], perto.numero_formatado)
        self.assertEqual(linhas[1][2], "12.345.678-9")
        self.assertEqual(linhas[1][4], "Autorização")
        self.assertEqual(linhas[1][8], "JANINE LACERDA DO PRADO")

    def test_linha_de_rascunho_vazio(self):
        self.oficio()
        r = self.lista()
        for texto in ["Sem roteiro", "Sem servidores", "Sem viatura", "Sem diárias", "Justificativa pendente"]:
            self.assertContains(r, texto)
        self.assertContains(r, 'class="st st--rascunho">Rascunho')

    def test_menu_unico_da_linha(self):
        o = self.oficio(dias=3, servidores=[self.janine], justificativa="texto")
        r = self.lista()
        self.assertContains(r, f"Ações do ofício {o.numero_formatado}")
        for texto in ["Abrir ofício", "Baixar documentos", "Anexar assinado", "Gere o PDF de um documento primeiro",
                      "Retificar ofício", "Ofício complementar", "Cancelar ofício", "Excluir ofício"]:
            self.assertContains(r, texto)
        self.assertContains(r, reverse("viagens_oficios:baixar", args=[o.pk]))
        # O modal lista ofício, justificativa e o termo de cada servidor.
        for valor in ["&quot;oficio&quot;", "&quot;justificativa&quot;", f"&quot;termo-{self.janine.pk}&quot;"]:
            self.assertContains(r, valor)
        self.assertContains(r, "data-baixar-dialogo")

    def test_leitor_nao_ve_menus_de_escrita(self):
        self.oficio(dias=3, servidores=[self.janine])
        self.user.groups.clear()
        r = self.lista()
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, "Ações do ofício")
        self.assertNotContains(r, "Novo ofício")

    def test_busca_por_numero_protocolo_motivo_e_destino(self):
        a = self.oficio(dias=3, protocolo="123456789")
        b = self.oficio()
        b.motivo = "Curso em Foz"
        b.save()
        self.assertContains(self.lista(q="12.345.678-9"), f"Nº {a.numero_formatado}")
        self.assertNotContains(self.lista(q="12.345.678-9"), f"Nº {b.numero_formatado}")
        self.assertContains(self.lista(q="Antonina"), f"Nº {a.numero_formatado}")
        self.assertContains(self.lista(q="Foz"), f"Nº {b.numero_formatado}")
        self.assertContains(self.lista(q=b.numero_formatado), f"Nº {b.numero_formatado}")
        self.assertNotContains(self.lista(q=b.numero_formatado), f"Nº {a.numero_formatado}")

    def test_sem_resultado(self):
        self.oficio()
        r = self.lista(q="ZZZ_PARIDADE")
        self.assertContains(r, "Nenhum ofício encontrado com os filtros aplicados.")

    def test_situacoes_na_trilha_com_contagem(self):
        futuro = self.oficio(dias=10)
        atual = self.oficio(dias=-3)
        rascunho = self.oficio()
        cancelado = self.oficio(dias=20, cancelar=True)
        contagens = {s["slug"]: s["total"] for s in self.lista().context["situacoes"]}
        # Sem data ainda não aconteceu: o rascunho conta entre os futuros.
        self.assertEqual(contagens, {"todas": 4, "futuras": 2, "atuais": 1, "finalizados": 0, "cancelados": 1})
        self.assertContains(self.lista(), "Contas prestadas")
        r = self.lista(situacao="atuais")
        self.assertEqual(r.context["situacao_ativa"], "atuais")
        self.assertContains(r, f"Nº {atual.numero_formatado}")
        self.assertNotContains(r, f"Nº {rascunho.numero_formatado}")
        self.assertNotContains(r, f"Nº {futuro.numero_formatado}")
        # A busca carrega a situação escolhida.
        self.assertContains(r, '<input type="hidden" name="situacao" value="atuais">')
        # Link antigo com várias situações continua filtrando.
        r = self.lista(situacao=["futuras", "cancelados"])
        self.assertContains(r, f"Nº {futuro.numero_formatado}")
        self.assertContains(r, f"Nº {cancelado.numero_formatado}")
        self.assertNotContains(r, f"Nº {atual.numero_formatado}")

    def test_paginacao_com_pagina(self):
        for _ in range(21):
            self.oficio()
        r = self.lista()
        self.assertContains(r, "Mostrando 1 a 20 de 21 ofícios")
        r = self.lista(pagina=2)
        self.assertContains(r, "Mostrando 21 a 21 de 21 ofícios")


class BaixarDocumentosTests(Cenario):
    def url(self, o):
        return reverse("viagens_oficios:baixar", args=[o.pk])

    def test_item_desconhecido_e_404(self):
        o = self.oficio(dias=3, servidores=[self.janine])
        r = self.client.post(self.url(o), {"itens": ["termo-999999"], "formato": "pdf"})
        self.assertEqual(r.status_code, 404)

    def test_sem_itens_volta_com_aviso(self):
        o = self.oficio(dias=3)
        volta = reverse("viagens_oficios:lista") + "?q=x"
        r = self.client.post(self.url(o), {"formato": "pdf", "next": volta})
        self.assertRedirects(r, volta)

    def test_cancelado_nao_baixa(self):
        o = self.oficio(dias=3, cancelar=True)
        r = self.client.post(self.url(o), {"itens": ["oficio"], "formato": "pdf"})
        self.assertRedirects(r, reverse("viagens_oficios:lista"))

    def test_leitor_nao_baixa(self):
        o = self.oficio(dias=3)
        self.user.groups.clear()
        r = self.client.post(self.url(o), {"itens": ["oficio"], "formato": "pdf"})
        self.assertEqual(r.status_code, 403)


class AcoesDaListaTests(Cenario):
    def test_novo_oficio_cria_rascunho_numerado_e_abre_o_cadastro(self):
        r = self.lista()
        self.assertContains(r, f'action="{reverse("viagens_oficios:criar")}"')
        r = self.client.post(reverse("viagens_oficios:criar"))
        o = Oficio.objects.get()
        self.assertRedirects(r, reverse("viagens_oficios:editar", args=[o.pk]))
        self.assertEqual(o.numero, 1)
        self.assertEqual(o.status, Oficio.STATUS_RASCUNHO)
        # Sem POST não se cria nada.
        self.assertRedirects(self.client.get(reverse("viagens_oficios:novo")), reverse("viagens_oficios:lista"))
        self.assertEqual(Oficio.objects.count(), 1)

    def test_novo_oficio_reaproveita_rascunho_vazio_abandonado(self):
        self.client.post(reverse("viagens_oficios:criar"))
        vazio = Oficio.objects.get()
        # Recém-aberto: pode estar sendo preenchido por alguém, então não se reaproveita.
        self.client.post(reverse("viagens_oficios:criar"))
        self.assertEqual(Oficio.objects.count(), 2)
        Oficio.objects.update(atualizado_em=timezone.now() - timedelta(hours=1))
        r = self.client.post(reverse("viagens_oficios:criar"))
        self.assertRedirects(r, reverse("viagens_oficios:editar", args=[vazio.pk]))
        self.assertEqual(Oficio.objects.count(), 2)
        # Com conteúdo, o rascunho é de alguém e não volta.
        Oficio.objects.update(protocolo="123456789")
        self.client.post(reverse("viagens_oficios:criar"))
        self.assertEqual(Oficio.objects.count(), 3)

    def test_acoes_voltam_para_a_lista_filtrada(self):
        o = self.oficio(dias=3)
        volta = reverse("viagens_oficios:lista") + "?situacao=futuras"
        r = self.client.post(reverse("viagens_oficios:acao", args=[o.pk, "retificar"]), {"next": volta})
        self.assertRedirects(r, volta)
        o.refresh_from_db()
        self.assertTrue(o.retificado_documento)
        # O mesmo item desliga a marca.
        self.client.post(reverse("viagens_oficios:acao", args=[o.pk, "retificar"]), {"next": volta})
        o.refresh_from_db()
        self.assertFalse(o.retificado_documento)
        r = self.client.post(reverse("viagens_oficios:acao", args=[o.pk, "excluir"]), {"next": volta})
        self.assertRedirects(r, volta)
        self.assertFalse(Oficio.objects.filter(pk=o.pk).exists())

    def test_cancelar_e_reativar_pela_lista(self):
        o = self.oficio(dias=3)
        self.client.post(reverse("viagens_oficios:acao", args=[o.pk, "cancelar"]), {"motivo": "Adiado"})
        o.refresh_from_db()
        self.assertTrue(o.cancelado)
        self.assertContains(self.lista(situacao="cancelados"), "Reativar ofício")
        self.client.post(reverse("viagens_oficios:acao", args=[o.pk, "reativar"]))
        o.refresh_from_db()
        self.assertFalse(o.cancelado)


class CadastroTests(Cenario):
    """O cadastro de ofício com os blocos e campos do Gerenciador de Viagens."""

    def payload(self, **extra):
        dados = {"numero": "", "protocolo": "12.345.678-9", "custeio": "UNIDADE_DPC", "motivo": "Missão",
                 "servidores": [str(self.janine.pk), str(self.joao.pk)],
                 "servidores_termo_autorizacao_present": "1",
                 "servidores_termo_autorizacao": [str(self.janine.pk)],
                 "viatura": self.duster.pk, "motorista_modo": "SERVIDOR", "motorista": self.janine.pk,
                 "acao": "rascunho"}
        dados.update(extra)
        return dados

    def editar(self, o):
        return self.client.get(reverse("viagens_oficios:editar", args=[o.pk]))

    def test_quatro_etapas_com_os_campos_da_origem(self):
        o = self.oficio(dias=3, servidores=[self.janine])
        r = self.editar(o)
        self.assertContains(r, f'Cadastro de ofício <span class="frm-ident">{o.numero_formatado}</span>')
        for texto in ["Dados e viajantes", "N° do Ofício", "Data do ofício", "Protocolo", "Custeio", "Nome da Instituição",
                      "Modelo de motivo", "Descrição", "Servidores", "Adicionar à equipe",
                      "buscar por nome, CPF ou RG", "Novo viajante", "Motorista", "Termo", "Viatura", "Escolher viatura",
                      "Buscar por placa ou modelo", "data-lista-escolha=\"viatura\"", "Motorista", "Condutor da viatura", "No sistema", "Manual",
                      "Buscar motorista no sistema", "Nome completo", "Ofício de origem",
                      "Roteiro e diárias", "Vincular a um roteiro existente", "Buscar roteiro existente", "Trechos", "Diárias",
                      "Justificativa", "Documentos", "data-de-embutir", "Novo viajante", "Nova viatura",
                      "Documento original (Ofício)", "Termos de Autorização"]:
            with self.subTest(texto=texto):
                self.assertContains(r, texto)
        # O que a origem não tem nesta tela fica de fora.
        for texto in ['name="assunto"', 'name="solicitante"', 'name="porte_transporte_armas"',
                      'name="transporte_placa_manual"', 'name="motorista_manual_cpf"', 'name="roteiro"',
                      "Regra de prazo", "Encerramento", "Faltam informações",
                      "Resumo do ofício", "Viatura e condução", "Equipe vinculada a este ofício"]:
            with self.subTest(ausente=texto):
                self.assertNotContains(r, texto)

    def test_rascunho_grava_e_volta_para_a_lista(self):
        o = self.oficio()
        volta = reverse("viagens_oficios:lista") + "?situacao=futuras"
        r = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(next=volta))
        self.assertRedirects(r, volta)
        o.refresh_from_db()
        self.assertEqual(o.protocolo, "123456789")
        self.assertEqual(set(o.servidores.all()), {self.janine, self.joao})
        self.assertEqual(list(o.servidores_termo_autorizacao.all()), [self.janine])
        self.assertEqual(o.viatura, self.duster)
        self.assertEqual(o.motorista, self.janine)
        self.assertEqual(o.status, Oficio.STATUS_RASCUNHO)

    def test_data_do_oficio_editavel(self):
        o = self.oficio()
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(data_criacao="2026-03-05"))
        o.refresh_from_db()
        self.assertEqual(str(o.data_criacao), "2026-03-05")
        # Em branco, a data gravada fica.
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(data_criacao=""))
        o.refresh_from_db()
        self.assertEqual(str(o.data_criacao), "2026-03-05")

    def test_numero_editavel_sem_repetir(self):
        a, b = self.oficio(), self.oficio()
        r = self.client.post(reverse("viagens_oficios:editar", args=[b.pk]), self.payload(numero=str(a.numero)))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, f"Já existe um ofício com o número {a.numero}")
        r = self.client.post(reverse("viagens_oficios:editar", args=[b.pk]), self.payload(numero="77"))
        self.assertEqual(r.status_code, 302)
        b.refresh_from_db()
        self.assertEqual(b.numero, 77)

    def test_motorista_de_fora_leva_o_oficio_de_origem(self):
        o = self.oficio()
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(
            servidores=[str(self.janine.pk)], motorista=self.joao.pk,
            motorista_oficio_referencia="15", motorista_protocolo_ref="11.222.333-4"))
        o.refresh_from_db()
        self.assertEqual(o.motorista, self.joao)
        self.assertEqual(o.motorista_oficio_referencia, f"15/{o.ano}")
        self.assertEqual(o.motorista_protocolo_ref, "112223334")
        # Motorista da equipe dispensa a referência.
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(
            motorista_oficio_referencia="15", motorista_protocolo_ref="11.222.333-4"))
        o.refresh_from_db()
        self.assertEqual(o.motorista_oficio_referencia, "")

    def test_motorista_manual(self):
        o = self.oficio()
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(
            motorista_modo="MANUAL", motorista_manual_nome="fulano de tal"))
        o.refresh_from_db()
        self.assertIsNone(o.motorista)
        self.assertEqual(o.motorista_manual_nome, "FULANO DE TAL")

    def test_campos_fora_da_tela_sao_preservados(self):
        o = self.oficio()
        Oficio.objects.filter(pk=o.pk).update(assunto="Interno", porte_transporte_armas=False)
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload())
        o.refresh_from_db()
        self.assertEqual(o.assunto, "Interno")
        self.assertFalse(o.porte_transporte_armas)

    def test_sem_sentinela_todos_tem_termo(self):
        o = self.oficio()
        dados = self.payload()
        dados.pop("servidores_termo_autorizacao_present")
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), dados)
        self.assertEqual(set(o.servidores_termo_autorizacao.all()), {self.janine, self.joao})

    def test_servidor_da_unidade_emissora_entra_sem_termo(self):
        from viagens_cadastros.models import ConfiguracaoSistema, Unidade
        cfg = ConfiguracaoSistema.para_usuario(self.user)
        cfg.unidade = self.ascom
        cfg.save()
        outra = Unidade.objects.create(nome="OUTRA UNIDADE", sigla="OUT")
        self.joao.unidade = outra
        self.joao.save()
        o = self.oficio(servidores=[self.janine, self.joao])
        o.servidores_termo_autorizacao.clear()
        pagina = self.client.get(reverse("viagens_oficios:editar", args=[o.pk])).content.decode()
        self.assertIn(f'data-unidade-emissora="{self.ascom.pk}"', pagina)
        # Sem termo escolhido ainda, a tela abre com termo só para quem é de fora.
        self.assertIn(f'name="servidores_termo_autorizacao" value="{self.joao.pk}" checked', pagina)
        self.assertNotIn(f'name="servidores_termo_autorizacao" value="{self.janine.pk}" checked', pagina)
        # Quem quiser ainda marca o termo à mão.
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(
            servidores_termo_autorizacao=[str(self.janine.pk), str(self.joao.pk)]))
        self.assertEqual(set(o.servidores_termo_autorizacao.all()), {self.janine, self.joao})

    def test_finalizar_com_pendencias_nao_finaliza(self):
        o = self.oficio()
        r = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(acao="finalizar"), follow=True)
        self.assertRedirects(r, reverse("viagens_oficios:editar", args=[o.pk]))
        # As pendências só aparecem depois de tentar finalizar.
        self.assertContains(r, "Associe um roteiro ao ofício.")
        o.refresh_from_db()
        self.assertEqual(o.status, Oficio.STATUS_RASCUNHO)
        self.assertContains(r, "Salvar rascunho")
        self.assertNotContains(r, "Finalizar Ofício")

    def test_finalizar_sem_pendencias(self):
        o = self.oficio(dias=20, protocolo="123456789", servidores=[self.janine], motorista=self.janine, viatura=self.duster)
        r = self.editar(o)
        self.assertContains(r, "Finalizar Ofício")
        r = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(
            servidores=[str(self.janine.pk)], servidores_termo_autorizacao=[str(self.janine.pk)], acao="finalizar"))
        self.assertRedirects(r, reverse("viagens_oficios:lista"))
        o.refresh_from_db()
        self.assertEqual(o.status, Oficio.STATUS_FINALIZADO)

    def test_cabecalho_mostra_autorizacao_ou_convalidacao_e_por_que(self):
        o = self.oficio(dias=5)
        r = self.editar(o)
        self.assertContains(r, "data-ofc-tipo>Autorização</span>")
        self.assertContains(r, "Justificativa obrigatória: 5 dias de antecedência, o prazo mínimo é de 10.")
        o = self.oficio(dias=-3)
        r = self.editar(o)
        self.assertContains(r, "data-ofc-tipo>Convalidação</span>")
        saida = self.hoje - timedelta(days=3)
        self.assertContains(r, f"A viagem começou em {saida:%d/%m/%Y}, antes da data do ofício ({self.hoje:%d/%m/%Y}).")
        o = self.oficio(dias=30)
        o.retificado_documento = True
        o.save()
        r = self.editar(o)
        self.assertContains(r, "data-ofc-tipo>Autorização · Retificado</span>")
        self.assertContains(r, "Justificativa dispensada: 30 dias de antecedência")
        r = self.editar(self.oficio())
        self.assertContains(r, "Sem data de saída no roteiro: por enquanto vale Autorização.")

    def test_condutores_autorizados_da_viatura(self):
        self.duster.motoristas.set([self.joao])
        o = self.oficio()
        r = self.editar(o)
        self.assertContains(r, f'data-motoristas="{self.joao.pk}"')
        # Motorista fora dos autorizados: grava, mas avisa.
        r = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(), follow=True)
        o.refresh_from_db()
        self.assertEqual(o.motorista, self.janine)
        self.assertContains(r, "não está entre os condutores autorizados da viatura AAA-1234")
        r = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(motorista=self.joao.pk), follow=True)
        self.assertNotContains(r, "condutores autorizados")
        # Viatura sem condutores cadastrados não tem restrição.
        self.duster.motoristas.clear()
        r = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(), follow=True)
        self.assertNotContains(r, "condutores autorizados")

    def test_oficio_do_motorista_de_fora_e_sugerido(self):
        from viagens_oficios.services import pendencias_motorista_documento
        # O ofício em que João viaja, com período que se sobrepõe.
        dele = self.oficio(dias=10, protocolo="112223334", servidores=[self.joao])
        antigo = self.oficio(servidores=[self.joao])
        self.oficio(dias=10, servidores=[self.joao], cancelar=True)
        o = self.oficio(dias=10, servidores=[self.janine])
        url = reverse("viagens_oficios:oficios_do_motorista", args=[o.pk])
        dados = self.client.get(url, {"motorista": self.joao.pk}).json()["oficios"]
        self.assertEqual([d["id"] for d in dados], [dele.pk, antigo.pk])
        self.assertEqual(dados[0]["referencia"], f"{dele.numero}/{dele.ano}")
        self.assertEqual(dados[0]["protocolo"], "11.222.333-4")
        self.assertEqual(dados[0]["motivo"], "Período que se sobrepõe")
        self.assertTrue(dados[0]["provavel"])
        self.assertFalse(dados[1]["provavel"])
        self.assertEqual(self.client.get(url, {"motorista": "x"}).json(), {"oficios": []})
        # O cartão sabe onde buscar.
        self.assertContains(self.editar(o), f'data-url="{url}"')
        # Ofício de número com 4 dígitos passa na conferência.
        o.motorista, o.motorista_oficio_referencia, o.motorista_protocolo_ref = self.joao, "1000/2026", "112223334"
        o.save()
        self.assertEqual(pendencias_motorista_documento(o), [])

    def test_campos_automaticos_nos_modelos(self):
        from viagens_oficios.campos_modelo import aplicar
        modelo = ModeloMotivoOficio.objects.create(nome="COBERTURA", texto="Cobertura em {destino} de {periodo}, com {servidores}. {desconhecido}")
        ModeloJustificativa.objects.create(nome="URGENTE", texto="Pedido feito {dias_antecedencia} dias antes (prazo {prazo}).")
        o = self.oficio(dias=5, servidores=[self.janine, self.joao])
        r = self.editar(o)
        saida = self.hoje + timedelta(days=5)
        volta = saida + timedelta(days=5)
        esperado = (f"Cobertura em ANTONINA/PR de {saida:%d/%m/%Y} a {volta:%d/%m/%Y}, "
                    "com JANINE LACERDA DO PRADO e JOÃO MARIO DE GOES. {desconhecido}")
        textos = r.context["modelos_texto"]
        self.assertEqual(textos["modelo_motivo"][modelo.pk], esperado)
        self.assertIn("Pedido feito 5 dias antes (prazo 10).", textos["justificativa-modelo"].values())
        # Sem roteiro, o marcador fica à vista e é preenchido na gravação.
        vazio = self.oficio(servidores=[self.janine])
        self.assertIn("{destino}", self.editar(vazio).context["modelos_texto"]["modelo_motivo"][modelo.pk])
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(motivo="Em {destino} por {evento}."))
        o.refresh_from_db()
        self.assertEqual(o.motivo, "Em ANTONINA/PR por {evento}.")
        # Chave solta não derruba nada.
        self.assertEqual(aplicar("texto { solto", {"destino": "X"}), "texto { solto")

    def _finalizar(self, o, **extra):
        return self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(
            servidores=[str(self.janine.pk)], servidores_termo_autorizacao=[str(self.janine.pk)], acao="finalizar", **extra),
            follow=True)

    def test_finalizar_confere_com_a_data_de_hoje(self):
        # Rascunho antigo: pela data dele a saída tinha folga; pela de hoje, não.
        o = self.oficio(dias=5, protocolo="123456789", servidores=[self.janine], motorista=self.janine,
                        viatura=self.duster, data_criacao=self.hoje - timedelta(days=20))
        r = self._finalizar(o)
        o.refresh_from_db()
        self.assertEqual(o.status, Oficio.STATUS_RASCUNHO)
        self.assertContains(r, "Informe o texto da justificativa.")
        # O rascunho continua com a data que tinha.
        self.assertEqual(o.data_criacao, self.hoje - timedelta(days=20))
        # Com a justificativa, finaliza com a data de hoje — e o PDF pode sair.
        self._finalizar(o, **{"justificativa-texto": "Convite recebido em cima da hora."})
        o.refresh_from_db()
        self.assertEqual(o.status, Oficio.STATUS_FINALIZADO)
        self.assertEqual(o.data_criacao, self.hoje)
        from viagens_oficios.services import validar_oficio_para_documento
        self.assertEqual(validar_oficio_para_documento(o)["pendencias"], [])

    def test_finalizar_mantem_a_data_digitada(self):
        o = self.oficio(dias=30, protocolo="123456789", servidores=[self.janine], motorista=self.janine, viatura=self.duster)
        digitada = self.hoje - timedelta(days=2)
        self._finalizar(o, data_criacao=digitada.isoformat())
        o.refresh_from_db()
        self.assertEqual(o.status, Oficio.STATUS_FINALIZADO)
        self.assertEqual(o.data_criacao, digitada)

    def test_finalizar_em_outro_ano_avisa(self):
        o = self.oficio(dias=30, protocolo="123456789", servidores=[self.janine], motorista=self.janine, viatura=self.duster)
        Oficio.objects.filter(pk=o.pk).update(ano=self.hoje.year - 1)
        o.refresh_from_db()
        r = self._finalizar(o)
        o.refresh_from_db()
        self.assertEqual(o.status, Oficio.STATUS_FINALIZADO)
        self.assertContains(r, "Confira se ele deve ser renumerado")

    def test_roteiro_montado_no_cadastro_fica_ligado_ao_oficio(self):
        o = self.oficio()
        saida = self.hoje + timedelta(days=15)
        dados = self.payload(**{
            "origem_municipio": self.curitiba.pk,
            "destinos-TOTAL_FORMS": "1", "destinos-INITIAL_FORMS": "0",
            "destinos-MIN_NUM_FORMS": "0", "destinos-MAX_NUM_FORMS": "1000",
            "destinos-0-municipio": self.antonina.pk, "destinos-0-ordem": "1",
            "trechos-TOTAL_FORMS": "1", "trechos-INITIAL_FORMS": "0",
            "trechos-MIN_NUM_FORMS": "0", "trechos-MAX_NUM_FORMS": "1000",
            "trechos-0-ordem": "1", "trechos-0-origem_municipio": self.curitiba.pk,
            "trechos-0-destino_municipio": self.antonina.pk,
            "trechos-0-saida_data": saida.isoformat(), "trechos-0-saida_hora": "08:00",
        })
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), dados)
        o.refresh_from_db()
        self.assertIsNotNone(o.roteiro)
        self.assertEqual(o.roteiro.origem_municipio, self.curitiba)
        self.assertEqual([d.municipio for d in o.roteiro.destinos.all()], [self.antonina])

    def test_vincular_roteiro_existente(self):
        o = self.oficio()
        r = self.roteiro(10)
        resposta = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(
            acao="vincular_roteiro", vincular_roteiro="1", roteiro_existente=str(r.pk), protocolo="98.765.432-1"))
        self.assertRedirects(resposta, reverse("viagens_oficios:editar", args=[o.pk]) + "#roteiro", fetch_redirect_response=False)
        o.refresh_from_db()
        self.assertEqual(o.roteiro, r)
        self.assertEqual(o.protocolo, "987654321")  # o que já estava digitado é gravado junto
        # Reaberto, o vínculo vem ligado: só a busca, sem sede e destinos à vista.
        pagina = self.client.get(reverse("viagens_oficios:editar", args=[o.pk])).content.decode()
        self.assertIn('name="vincular_roteiro" value="1"', pagina)
        self.assertIn('data-roteiro-proprio hidden', pagina)
        self.assertIn(f'name="roteiro_existente" value="{r.pk}" checked', pagina)

    def test_protocolo_repetido_avisa_sem_impedir(self):
        outro = self.oficio(protocolo="123456789")
        o = self.oficio(dias=5)
        r = self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), self.payload(protocolo="12.345.678-9"), follow=True)
        o.refresh_from_db()
        self.assertEqual(o.protocolo, "123456789")
        self.assertContains(r, f"também está no ofício {outro.numero_formatado}")

    def test_previa_usa_o_tamanho_da_equipe(self):
        from unittest import mock
        with mock.patch("viagens_roteiros.views.previa_diarias") as previa:
            previa.return_value = {"trechos": [], "totais": {
                "total_valor": "1,00", "resumo_diarias": "", "valor_extenso": "", "quantidade_servidores": 3, "valor_por_servidor": ""}}
            self.client.post(reverse("viagens_roteiros:previa_diarias"), {"diarias_servidores": "3"})
        self.assertEqual(previa.call_args.args[2], 3)

    def test_roteiro_trocado_na_tela_e_so_vinculado(self):
        o = self.oficio(dias=5)
        antigo = o.roteiro
        novo = self.roteiro(12)
        trechos_novo = novo.trechos.count()
        # O editor chega com o percurso do escolhido como linhas novas (sem id).
        dados = self.payload(vincular_roteiro="1", roteiro_trocado="1", roteiro_existente=str(novo.pk), **{
            "origem_municipio": self.curitiba.pk,
            "destinos-TOTAL_FORMS": "1", "destinos-INITIAL_FORMS": "0",
            "destinos-0-municipio": self.antonina.pk, "destinos-0-ordem": "1",
            "trechos-TOTAL_FORMS": "1", "trechos-INITIAL_FORMS": "0",
            "trechos-0-ordem": "1", "trechos-0-origem_municipio": self.curitiba.pk,
            "trechos-0-destino_municipio": self.antonina.pk,
        })
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), dados)
        o.refresh_from_db()
        self.assertEqual(o.roteiro, novo)
        self.assertEqual(novo.trechos.count(), trechos_novo)  # nada duplicado no escolhido
        self.assertTrue(type(antigo).objects.filter(pk=antigo.pk).exists())

    def test_dados_do_roteiro_trazem_trechos_e_rota(self):
        r = self.roteiro(7)
        dados = self.client.get(reverse("viagens_roteiros:dados", args=[r.pk])).json()
        self.assertEqual(len(dados["trechos"]), 1)
        self.assertEqual(dados["trechos"][0]["saida_hora"], "08:00")
        self.assertIn("rota", dados)

    def test_sem_roteiro_o_vinculo_vem_desligado(self):
        o = self.oficio()
        pagina = self.client.get(reverse("viagens_oficios:editar", args=[o.pk])).content.decode()
        self.assertIn('name="vincular_roteiro" value=""', pagina)
        self.assertIn('data-roteiro-vinculo hidden', pagina)
        self.assertNotIn('data-roteiro-proprio hidden', pagina)

    def test_vinculo_ligado_sem_escolha_nao_mexe_no_roteiro(self):
        o = self.oficio()
        dados = self.payload(vincular_roteiro="1", **{"origem_municipio": self.curitiba.pk,
                             "trechos-TOTAL_FORMS": "0", "trechos-INITIAL_FORMS": "0",
                             "destinos-TOTAL_FORMS": "1", "destinos-INITIAL_FORMS": "0",
                             "destinos-0-municipio": self.antonina.pk})
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), dados)
        o.refresh_from_db()
        self.assertIsNone(o.roteiro)

    def test_editor_vazio_nao_cria_roteiro(self):
        o = self.oficio()
        dados = self.payload(**{"origem_municipio": "", "trechos-TOTAL_FORMS": "0", "trechos-INITIAL_FORMS": "0",
                                "destinos-TOTAL_FORMS": "1", "destinos-INITIAL_FORMS": "0", "destinos-0-municipio": ""})
        self.client.post(reverse("viagens_oficios:editar", args=[o.pk]), dados)
        o.refresh_from_db()
        self.assertIsNone(o.roteiro)

    def test_conferencia(self):
        o = self.oficio(dias=-20, protocolo="123456789", servidores=[self.janine, self.joao], motorista=self.joao, viatura=self.duster, justificativa="teste 1")
        r = self.editar(o)
        for texto in ["Termo de Autorização — JANINE LACERDA DO PRADO", "Visualizar documento", "Baixar PDF", "Baixar PDFs",
                      reverse("viagens_oficios:visualizar", args=[o.pk, "oficio"]),
                      reverse("viagens_oficios:visualizar_termo", args=[o.pk, self.janine.pk]),
                      reverse("documentos:editor_embutido", args=["oficio", o.pk]),
                      reverse("documentos:editor_embutido", args=["termo_oficio", o.pk]) + f"?v={self.janine.pk}"]:
            with self.subTest(texto=texto):
                self.assertContains(r, texto)

    def test_documentos_indisponiveis_com_pendencia(self):
        o = self.oficio()
        r = self.editar(o)
        # Sem o PDF para visualizar; o cartão abre o editor, que mostra as pendências.
        self.assertNotContains(r, reverse("viagens_oficios:visualizar", args=[o.pk, "oficio"]))
        self.assertContains(r, reverse("documentos:editor_embutido", args=["oficio", o.pk]))

    def test_acao_volta_para_o_formulario(self):
        o = self.oficio(dias=3, servidores=[self.janine])
        r = self.client.post(reverse("viagens_oficios:acao", args=[o.pk, "arquivar"]))
        self.assertRedirects(r, reverse("viagens_oficios:editar", args=[o.pk]))

    def test_formulario_abre_sem_aviso(self):
        # Tela de cadastro não abre com aviso: nem pendências, nem "pronto para emissão".
        for o in (self.oficio(), self.oficio(dias=3, protocolo="123456789", servidores=[self.janine], viatura=self.duster, motorista=self.janine)):
            r = self.editar(o)
            for texto in ["Faltam informações para emitir o documento", "Completar o ofício", "pronto para emissão", "aviso--callout"]:
                self.assertNotContains(r, texto)

    def test_leitor_nao_abre_o_cadastro(self):
        o = self.oficio()
        self.user.groups.clear()
        self.assertEqual(self.editar(o).status_code, 403)
        self.assertEqual(self.client.get(reverse("viagens_oficios:visualizar", args=[o.pk, "oficio"])).status_code, 403)


class CatalogosTests(Cenario):
    def test_catalogos_no_padrao_dos_cadastros(self):
        ModeloMotivoOficio.objects.create(nome="COBERTURA", texto="Texto", is_padrao=True)
        ModeloMotivoOficio.objects.create(nome="CAPACITAÇÃO", texto="Texto")
        r = self.client.get(reverse("viagens_oficios:catalogo", args=["motivos"]))
        self.assertEqual(r.status_code, 200)
        self.assertTemplateUsed(r, "pages/viagens_cadastros/lista.html")
        for texto in ["Motivos de ofício", "COBERTURA", "CAPACITAÇÃO", "Definir padrão", "Excluir", "Novo modelo de motivo"]:
            self.assertContains(r, texto)
        r = self.client.get(reverse("viagens_oficios:catalogo_novo", args=["motivos"]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Texto que vai para o ofício ao escolher este modelo")
        r = self.client.get(reverse("viagens_oficios:catalogo", args=["justificativas"]))
        self.assertContains(r, "Modelos de justificativa")

    def test_criar_e_definir_padrao(self):
        r = self.client.post(reverse("viagens_cadastros:novo", args=["modelos-justificativa"]),
                             {"nome": "demanda urgente", "texto": "Texto", "ordem": 10, "ativo": "1", "ativo__enviado": "1", "is_padrao__enviado": "1"})
        self.assertEqual(r.status_code, 302)
        m = ModeloJustificativa.objects.get()
        self.assertEqual(m.nome, "DEMANDA URGENTE")
        self.assertFalse(m.is_padrao)
        self.client.post(reverse("viagens_cadastros:definir_padrao", args=["modelos-justificativa", m.pk]))
        m.refresh_from_db()
        self.assertTrue(m.is_padrao)
        # Fora da trilha de cadastros (o rótulo segue na gaveta "Modelos" da navegação).
        self.assertNotContains(self.client.get(reverse("viagens_cadastros:index"), follow=True), 'class="cad-i__t">Motivos de ofício')
