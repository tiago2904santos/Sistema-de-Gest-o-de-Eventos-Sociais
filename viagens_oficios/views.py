import io
from zipfile import ZipFile

from django.contrib import messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction, IntegrityError
from django.http import Http404, HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.views.decorators.http import require_GET, require_http_methods, require_POST

from core.listagens import ITENS_POR_PAGINA, paginar, opcoes_choices
from core.numeracao import bloquear_escopo_numeracao, NAMESPACE_OFICIO
from core.utils.masks import format_protocolo
from viagens_cadastros.permissions import acesso_ao_modulo, pode_editar_cadastros, eh_gestor_viagens
from documentos.services.types import DocumentoTipo, DocumentoFormato
from documentos.services.responses import build_download_response, build_inline_pdf_response
from documentos.services.exceptions import DocumentError
from .models import Oficio, ConfiguracaoNumeracaoOficio, ModeloMotivoOficio, ModeloJustificativa
from .forms import OficioForm, JustificativaForm, NumeracaoForm
from .selectors import listar_oficios, get_oficio_by_id
from .services import reservar_numero_oficio, excluir_oficio, retificar_oficio, marcar_oficio_complementar, validar_oficio_para_documento
from .justificativas_services import atualizar_justificativa_oficio, avaliar_justificativa_oficio
from .document_generation import gerar_documento
from .view_helpers import campos_v32


def exigir_operador(request):
    if not pode_editar_cadastros(request.user):
        raise PermissionDenied


@acesso_ao_modulo
def preview_artefato(request, pk):
    from documentos.services.access import obter_artefato_para_download
    from django.http import FileResponse
    artefato = obter_artefato_para_download(request.user, pk)
    if artefato.formato != 'pdf':
        raise Http404
    response = FileResponse(artefato.arquivo_efetivo.open('rb'), content_type='application/pdf')
    response['Cache-Control'] = 'no-store'
    return response


@acesso_ao_modulo
@require_http_methods(['GET', 'POST'])
def assinatura_artefato(request, pk):
    from django import forms
    from documentos.services.access import obter_artefato_para_download
    from documentos.services.persistence import anexar_arquivo_assinado, remover_arquivo_assinado
    from core.retorno import next_valido, voltar_para
    exigir_operador(request)
    artefato = obter_artefato_para_download(request.user, pk)
    if artefato.formato != 'pdf' or not (artefato.oficio_id or artefato.termo_id):
        raise Http404
    voltar = reverse('viagens_termos:editar', args=[artefato.termo_id]) if artefato.termo_id else reverse('viagens_oficios:editar', args=[artefato.oficio_id])

    class UploadForm(forms.Form):
        arquivo = forms.FileField(label='Documento assinado (PDF)')

    form = UploadForm(request.POST or None, request.FILES or None)
    # Pelo modal, o POST traz `next` (a página de onde se abriu): o retorno vai
    # para lá, e um erro volta como mensagem em vez de abrir esta página.
    do_modal = request.method == 'POST' and bool(next_valido(request))
    retorno = voltar_para(request, voltar)
    if request.method == 'POST':
        try:
            if request.POST.get('acao') == 'remover':
                remover_arquivo_assinado(artefato)
                messages.success(request, 'Versão assinada removida. O PDF gerado volta a valer.')
                return redirect(retorno)
            if form.is_valid():
                anexar_arquivo_assinado(artefato, form.cleaned_data['arquivo'])
                messages.success(request, 'Documento assinado anexado. A versão anterior permanece no histórico.')
                return redirect(retorno)
        except DocumentError as exc:
            form.add_error('arquivo', str(exc))
        if do_modal:
            messages.error(request, ' '.join(form.errors.get('arquivo', [])) or 'Não foi possível anexar o documento.')
            return redirect(retorno)
    return render(request, 'pages/viagens_oficios/assinatura.html', {'form': form, 'artefato': artefato, 'url_voltar': voltar})


def _artefatos_pdf_da_pagina(oficios):
    """Último PDF de cada documento dos ofícios da página, para "Visualizar" e "Anexar assinado".

    Uma consulta para a página inteira: (ofício, tipo, servidor) → artefato.
    """
    from documentos.models import DocumentoArtefato
    ids = [o.pk for o in oficios]
    if not ids:
        return {}
    mapa = {}
    for art in DocumentoArtefato.objects.filter(oficio_id__in=ids, formato='pdf').order_by('criado_em').values_list('oficio_id', 'tipo', 'servidor_id', 'pk'):
        mapa[(art[0], art[1], art[2])] = art[3]
    return mapa


#: Os filtros prontos de `listar_oficios` que a lista oferece, com o rótulo do chip.
FILTROS_DA_LISTA = (
    ('viagem_de', 'Viagem a partir de'),
    ('viagem_ate', 'Viagem até'),
    ('criacao_de', 'Ofício a partir de'),
    ('criacao_ate', 'Ofício até'),
    ('ano', 'Ano'),
)


def _filtros_da_lista(request):
    """Busca, situações, períodos, ano e ordenação pedidos na URL da lista."""
    from datetime import date
    from . import abas as abas_de_oficio
    from .selectors import normalizar_ordenacao

    filtros = {}
    for nome, _rotulo in FILTROS_DA_LISTA:
        valor = request.GET.get(nome, '').strip()
        if nome == 'ano':
            valor = valor if valor.isdigit() and len(valor) == 4 else ''
        elif valor:
            try:
                date.fromisoformat(valor)
            except ValueError:
                valor = ''
        filtros[nome] = valor
    sort = request.GET.get('sort', '')
    return {
        'q': request.GET.get('q', '').strip(),
        'situacoes': abas_de_oficio.normalizar_abas(request.GET.getlist('situacao')),
        'filtros': filtros,
        'sort': normalizar_ordenacao(sort) if sort else '',
    }


def _oficios_filtrados(pedido, *, com_situacoes=True):
    return listar_oficios(
        pedido['q'], ano=pedido['filtros']['ano'],
        situacoes=pedido['situacoes'] if com_situacoes else None,
        viagem_de=pedido['filtros']['viagem_de'], viagem_ate=pedido['filtros']['viagem_ate'],
        criacao_de=pedido['filtros']['criacao_de'], criacao_ate=pedido['filtros']['criacao_ate'],
        sort=pedido['sort'],
    )


@acesso_ao_modulo
def lista(request):
    """Lista de ofícios no padrão das listas de termos, roteiros e
    justificativas: situações na trilha à esquerda, busca na barra do cartão,
    uma célula por ofício e um menu único de ações. Os filtros prontos
    (períodos, ano, ordenação) ficam em "Filtros", e o recorte sai em planilha."""
    from django.core.paginator import Paginator
    from core.retorno import daqui
    from datetime import date
    from . import abas as abas_de_oficio
    from .justificativas_services import get_prazo_justificativa_dias
    from .presenters import artefatos_pdf_por_oficio, linha_da_lista
    from .selectors import opcoes_de_ordenacao
    from viagens_prestacoes.importacao.entrada import limite_de_bytes

    pedido = _filtros_da_lista(request)
    q, escolhidas = pedido['q'], pedido['situacoes']
    base = _oficios_filtrados(pedido, com_situacoes=False)
    queryset = _oficios_filtrados(pedido)
    paginator = Paginator(queryset, ITENS_POR_PAGINA)
    pagina = paginator.get_page(request.GET.get('pagina'))
    parametros = request.GET.copy()
    parametros.pop('pagina', None)
    artefatos = artefatos_pdf_por_oficio(pagina.object_list)
    prazo = get_prazo_justificativa_dias()
    linhas = [linha_da_lista(o, artefatos_pdf=artefatos.get(o.pk, {}), prazo=prazo) for o in pagina]

    def url_da_situacao(aba=None):
        destino = parametros.copy()
        destino.pop('situacao', None)
        if aba:
            destino['situacao'] = aba
        return '?' + destino.urlencode()

    def url_sem(nome):
        destino = parametros.copy()
        destino.pop(nome, None)
        return '?' + destino.urlencode()

    icones = {
        abas_de_oficio.ABA_FUTURAS: 'calendar',
        abas_de_oficio.ABA_ATUAIS: 'clock',
        abas_de_oficio.ABA_FINALIZADOS: 'check-circle',
        abas_de_oficio.ABA_CANCELADOS: 'ban',
    }
    contagem = abas_de_oficio.contar_por_aba(base)
    situacoes = [{'slug': 'todas', 'titulo': 'Todos', 'total': base.count(), 'icone': 'checklist', 'url': url_da_situacao()}] + [
        {'slug': chave, 'titulo': rotulo, 'total': contagem[chave], 'icone': icones[chave], 'url': url_da_situacao(chave)}
        for chave, rotulo in abas_de_oficio.ABA_ROTULOS
    ]
    filtros = pedido['filtros']
    ativos = []
    for nome, rotulo in FILTROS_DA_LISTA:
        valor = filtros[nome]
        if valor:
            texto = valor if nome == 'ano' else date.fromisoformat(valor).strftime('%d/%m/%Y')
            ativos.append({'rotulo': f'{rotulo}: {texto}', 'url_remover': url_sem(nome)})
    ordenacoes = opcoes_de_ordenacao()
    if pedido['sort']:
        rotulo = next(o['rotulo'] for o in ordenacoes if o['valor'] == pedido['sort'])
        ativos.append({'rotulo': f'Ordem: {rotulo}', 'url_remover': url_sem('sort')})
    anos = sorted({a for a in Oficio.objects.exclude(ano__isnull=True).values_list('ano', flat=True)}, reverse=True)

    return render(request, 'pages/viagens_oficios/lista.html', {
        'linhas': linhas, 'pagina': pagina, 'querystring': parametros.urlencode(), 'q': q,
        'paginas_visiveis': list(paginator.get_elided_page_range(pagina.number, on_each_side=1, on_ends=1)),
        'elipse': paginator.ELLIPSIS,
        'situacoes': situacoes,
        'situacoes_escolhidas': escolhidas,
        'situacao_ativa': 'todas' if not escolhidas else escolhidas[0] if len(escolhidas) == 1 else '',
        'filtros': filtros, 'filtros_ativos': ativos, 'sort': pedido['sort'],
        'opcoes_ordenacao': ordenacoes,
        'opcoes_ano': [{'valor': str(a), 'rotulo': str(a)} for a in anos],
        'url_exportar': reverse('viagens_oficios:exportar') + ('?' + parametros.urlencode() if parametros else ''),
        'tem_filtros': bool(q or escolhidas or ativos), 'url_atual': daqui(request),
        'pode_editar': pode_editar_cadastros(request.user), 'gestor': eh_gestor_viagens(request.user),
        # O modal "Importar processo do eProtocolo" (importador da prestação) diz o limite do arquivo.
        'importacao_limite_mb': limite_de_bytes() // (1024 * 1024),
    })


@acesso_ao_modulo
@require_GET
def exportar(request):
    """O recorte da lista (busca, situações, filtros e ordem) numa planilha Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.utils import timezone
    from .justificativas_services import get_prazo_justificativa_dias
    from .presenters import destinos_resumidos, justificativa_do_cartao, tipo_do_oficio, transporte_do_cartao
    from .roteiro_context import periodo_roteiro

    pedido = _filtros_da_lista(request)
    prazo = get_prazo_justificativa_dias()
    colunas = ['Nº', 'Data do ofício', 'Protocolo', 'Situação', 'Tipo', 'Destinos', 'Saída', 'Retorno',
               'Servidores', 'Motorista', 'Viatura', 'Diárias (R$)', 'Quantidade de diárias', 'Justificativa']
    livro = Workbook()
    aba = livro.active
    aba.title = 'Ofícios'
    aba.append(colunas)
    for celula in aba[1]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill('solid', fgColor='D1D3D4')
    for o in _oficios_filtrados(pedido):
        saida, retorno = periodo_roteiro(o.roteiro) if o.roteiro_id else (None, None)
        transporte = transporte_do_cartao(o)
        try:
            diarias = o.diarias_para_servidores() if o.roteiro_id else None
        except ValidationError:  # roteiro sem efetivo: a linha sai sem o valor
            diarias = None
        motorista = o.motorista.nome if o.motorista_id else (o.motorista_manual_nome or '')
        tipo = tipo_do_oficio(o, prazo=prazo)
        aba.append([
            o.numero_formatado,
            o.data_criacao,
            format_protocolo(o.protocolo),
            'Cancelado' if o.cancelado else o.get_status_display(),
            tipo['rotulo'] + (f" ({tipo['marcador']})" if tipo['marcador'] else ''),
            destinos_resumidos(o.roteiro, maximo=100) if o.roteiro_id else '',
            timezone.localtime(saida).replace(tzinfo=None) if saida else None,
            timezone.localtime(retorno).replace(tzinfo=None) if retorno else None,
            ', '.join(s.nome for s in o.servidores.all()),
            motorista,
            ' · '.join(p for p in [transporte['modelo'], transporte['placa']] if p),
            diarias['valor_decimal'] if diarias else None,
            (diarias or {}).get('quantidade') or '',
            'Preenchida' if justificativa_do_cartao(o)['preenchida'] else ('Pendente' if tipo['justificativa_obrigatoria'] else ''),
        ])
    for linha in aba.iter_rows(min_row=2):
        linha[1].number_format = 'DD/MM/YYYY'
        linha[6].number_format = linha[7].number_format = 'DD/MM/YYYY HH:MM'
        linha[11].number_format = '#,##0.00'
    for indice, coluna in enumerate(colunas, start=1):
        maior = max([len(coluna)] + [len(str(c.value or '')) for c in aba[get_column_letter(indice)][1:]])
        aba.column_dimensions[get_column_letter(indice)].width = min(60, maior + 2)
    aba.freeze_panes = 'A2'
    buffer = io.BytesIO()
    livro.save(buffer)
    response = HttpResponse(buffer.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="oficios-{timezone.localdate():%Y-%m-%d}.xlsx"'
    response['Cache-Control'] = 'no-store'
    return response


@acesso_ao_modulo
@require_POST
def baixar(request, pk):
    """Os documentos marcados no modal "Baixar documentos" da lista.

    `itens`: `oficio`, `justificativa` e `termo-<servidor>`. `formato`: pdf ou
    docx. `saida`: `separados` (um arquivo, ou ZIP) ou `unico` (um PDF só, na
    ordem da lista). `versao`: `assinado` (padrão) ou `original`.
    """
    from core.retorno import voltar_para
    from viagens_termos.services import gerar_termo_um
    from viagens_termos.views import resposta_pdf_consolidado

    exigir_operador(request)
    oficio = get_oficio_by_id(pk)
    retorno = voltar_para(request, reverse('viagens_oficios:lista'))
    if oficio.cancelado:
        messages.error(request, 'Reative o ofício antes de baixar documentos.')
        return redirect(retorno)
    formato = request.POST.get('formato', 'pdf')
    if formato not in ('pdf', 'docx'):
        raise Http404
    servidores = {f'termo-{s.pk}': s for s in oficio.servidores_termo_autorizacao.all()}
    ordem = ['oficio', 'justificativa', *servidores]
    marcados = request.POST.getlist('itens')
    if set(marcados) - set(ordem):
        raise Http404
    pedidos = [v for v in ordem if v in marcados]
    if not pedidos:
        messages.error(request, 'Marque ao menos um documento para baixar.')
        return redirect(retorno)
    fmt = DocumentoFormato(formato)
    usar_assinado = request.POST.get('versao', 'assinado') != 'original'
    try:
        documentos = []
        for valor in pedidos:
            if valor in servidores:
                documentos.append(gerar_termo_um(oficio, servidores[valor], fmt, usar_assinado=usar_assinado))
            else:
                tipo = DocumentoTipo.OFICIO if valor == 'oficio' else DocumentoTipo.JUSTIFICATIVA
                documentos.append(gerar_documento(oficio, fmt, tipo, usar_assinado=usar_assinado))
    except (ValidationError, DocumentError) as exc:
        messages.error(request, '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc))
        return redirect(retorno)
    referencia = oficio.numero_formatado.replace('/', '-')
    if len(documentos) == 1:
        return resposta_documento(request, documentos[0])
    if fmt == DocumentoFormato.PDF and request.POST.get('saida') == 'unico':
        return resposta_pdf_consolidado(documentos, f'oficio-{referencia}-documentos.pdf')
    return resposta_lote(documentos, f'oficio-{referencia}-documentos.zip')


@acesso_ao_modulo
@require_POST
def criar(request):
    """"Novo ofício" da origem: cria o rascunho já numerado e abre o cadastro."""
    from core.retorno import com_next, next_valido
    from .services import criar_oficio_rascunho
    from viagens_viagem.services import viagem_do_request
    exigir_operador(request)
    oficio = criar_oficio_rascunho(viagem=viagem_do_request(request))
    destino = reverse('viagens_oficios:editar', args=[oficio.pk])
    retorno = next_valido(request)
    return redirect(com_next(destino, retorno) if retorno else destino)


def _url_de_volta(oficio, viagem_id=None):
    """Para onde se volta ao sair do ofício: a etapa 3 da viagem dele, ou a lista."""
    viagem_id = viagem_id or getattr(oficio, 'viagem_id', None)
    if viagem_id:
        return reverse('viagens_viagem:etapa', args=[viagem_id, 3])
    return reverse('viagens_oficios:lista')


CAMPOS_FORA_DO_HISTORICO = {'atualizado_em', 'criado_em'}


def historico_do_oficio(oficio, limite=20):
    """Os últimos registros da trilha sobre este ofício — o ofício em si e os
    blocos documentais dele (texto de parágrafo alterado, quebra de página).
    Cada registro sai com a origem e os campos que mudaram, para a tela."""
    from django.db.models import Q
    from auditoria.models import RegistroAuditoria
    blocos = list(oficio.blocos_documentais.values_list('pk', flat=True)) if oficio.pk else []
    filtro = Q(modelo='viagens_oficios.oficio', objeto_id=str(oficio.pk))
    if blocos:
        filtro |= Q(modelo='documentos.documentobloco', objeto_id__in=[str(pk) for pk in blocos])
    registros = list(RegistroAuditoria.objects.filter(filtro).select_related('usuario').order_by('-criado_em')[:limite])
    for registro in registros:
        registro.campos_alterados = sorted(
            campo for campo in registro.alteracoes if campo not in CAMPOS_FORA_DO_HISTORICO and campo not in ('novo', 'antigo')
        )
        registro.sobre_bloco = registro.modelo == 'documentos.documentobloco'
    return registros


def _editor_vazio(dados):
    """O editor de roteiro veio sem nada: sem sede, sem destino e sem trecho."""
    if dados.get('origem_municipio'):
        return False
    for chave, valor in dados.items():
        if valor and (chave.endswith('-municipio') or chave.endswith('-destino_municipio')):
            return False
    return True


def _gravar_roteiro(request, oficio, *, finalizar):
    """O editor de roteiro embutido, gravado como a tela de roteiros grava.

    Sem o editor no POST, ou com ele vazio num ofício ainda sem roteiro, não
    há o que gravar. O roteiro novo fica ligado ao ofício.
    """
    from viagens_roteiros.models import Roteiro
    from viagens_roteiros.views import gravar_editor
    dados = request.POST
    if 'trechos-TOTAL_FORMS' not in dados:
        return None
    roteiro = oficio.roteiro
    if roteiro is None and _editor_vazio(dados):
        return None
    rascunho = not finalizar and (roteiro is None or roteiro.status == Roteiro.Status.RASCUNHO)
    gravacao = gravar_editor(dados, roteiro, rascunho=rascunho, usuario=request.user)
    if gravacao.roteiro is not None and gravacao.roteiro.pk and oficio.roteiro_id != gravacao.roteiro.pk:
        oficio.roteiro = gravacao.roteiro
        oficio.save(update_fields=['roteiro', 'atualizado_em'])
    return gravacao


def _contexto_roteiro(request, oficio, gravacao):
    from viagens_roteiros.forms import DestinoFormSet, RoteiroForm, TrechoFormSet
    from viagens_roteiros.views import _contexto_do_form, _sede_inicial
    if gravacao is not None:
        roteiro = gravacao.roteiro if gravacao.roteiro is not None and gravacao.roteiro.pk else oficio.roteiro
        contexto = _contexto_do_form(roteiro, gravacao.form, gravacao.formset, gravacao.destinos)
    else:
        roteiro = oficio.roteiro
        form = RoteiroForm(instance=roteiro, initial=_sede_inicial(request, roteiro))
        contexto = _contexto_do_form(roteiro, form, TrechoFormSet(instance=roteiro), DestinoFormSet(instance=roteiro))
    # O ofício com roteiro abre com o vínculo ligado; a busca inclui o atual.
    if request.method == 'POST':
        contexto['vinculado'] = request.POST.get('vincular_roteiro') == '1'
    else:
        contexto['vinculado'] = bool(oficio.roteiro_id)
    contexto['opcoes_vinculo'] = _opcoes_vinculo(oficio.roteiro_id)
    contexto['vinculo_atual'] = str(oficio.roteiro_id or '')
    # As diárias da tela são as do ofício: valor por servidor × equipe.
    contexto['diarias_servidores'] = max(1, oficio.servidores.count())
    contexto['diarias_tela'] = _diarias_do_oficio(oficio)
    return contexto


def _diarias_do_oficio(oficio):
    from django.core.exceptions import ValidationError
    from viagens_roteiros.services.diarias import formatar_valor
    try:
        diarias = oficio.diarias_para_servidores()
    except ValidationError:
        return None
    if not diarias:
        return None
    return {
        'valor': formatar_valor(diarias['valor_decimal']),
        'extenso': diarias['valor_extenso'],
        'resumo': diarias['quantidade'],
        'servidores': diarias['quantidade_servidores'],
    }


def _opcoes_vinculo(atual=None):
    """Roteiros para vincular, como as linhas da lista de escolha: o percurso
    no nome e, embaixo, o período e as diárias. O vinculado vem primeiro."""
    from viagens_roteiros.models import Roteiro
    from viagens_roteiros.presenters import periodo_display
    opcoes = []
    roteiros = (Roteiro.objects.select_related("origem_municipio")
                .prefetch_related("destinos__municipio", "trechos").filter(cancelado=False).order_by("-atualizado_em")[:200])
    for r in roteiros:
        destinos = [d.municipio.nome for d in r.destinos.all() if d.municipio_id]
        sede = r.origem_municipio.nome if r.origem_municipio_id else "Sem sede"
        periodo = periodo_display(r)
        detalhes = [f"#{r.pk}"]
        if periodo and periodo != "—":
            detalhes.append(periodo)
        if r.resumo_diarias:
            detalhes.append(r.resumo_diarias)
        opcoes.append({
            "valor": str(r.pk),
            "rotulo": " → ".join([sede, *destinos]) if destinos else f"{sede} → sem destinos",
            "detalhes": " · ".join(detalhes),
            "busca": str(r.pk),
        })
    opcoes.sort(key=lambda o: o["valor"] != str(atual))
    return opcoes


def _vincular_roteiro(request, oficio):
    """O roteiro existente escolhido na busca passa a ser o do ofício."""
    from viagens_roteiros.models import Roteiro
    valor = request.POST.get('roteiro_existente', '')
    roteiro = Roteiro.objects.filter(pk=valor, cancelado=False).first() if valor.isdigit() else None
    if roteiro is None:
        return None
    if oficio.roteiro_id != roteiro.pk:
        oficio.roteiro = roteiro
        oficio.save(update_fields=['roteiro', 'atualizado_em'])
    return roteiro


def _aviso_de_condutor(oficio):
    """Avisa, sem impedir, quando o motorista não é condutor autorizado da viatura.

    Só vale para viatura com condutores cadastrados: lista vazia é viatura
    sem restrição.
    """
    if not (oficio.viatura_id and oficio.motorista_id):
        return ''
    autorizados = list(oficio.viatura.motoristas.all())
    if not autorizados or any(m.pk == oficio.motorista_id for m in autorizados):
        return ''
    return (f'{oficio.motorista.nome} não está entre os condutores autorizados da viatura '
            f'{oficio.viatura.placa_formatada}. Confira antes de emitir o ofício.')


def _conflitos_da_tela(oficio, form):
    """Conflitos de agenda do que está na tela (core/conflitos.py).

    Com o formulário recusado, valem as escolhas enviadas; senão, as gravadas.
    O período é sempre o do roteiro gravado.
    """
    from core.conflitos import conflitos_do_oficio
    if form.is_bound:
        dados = form.data
        manual = dados.get('motorista_modo') == Oficio.MOTORISTA_MODO_MANUAL
        return conflitos_do_oficio(
            oficio, servidores=dados.getlist('servidores'), viatura=dados.get('viatura'),
            motorista=None if manual else dados.get('motorista'),
        )
    return conflitos_do_oficio(oficio)


def _avisos_de_conflito(oficio, limite=5):
    """Os conflitos em texto para o aviso depois de salvar (sem bloquear)."""
    from core.conflitos import conflitos_do_oficio
    achados = conflitos_do_oficio(oficio)
    textos = [f'Conflito de agenda: {c.mensagem}.' for c in achados[:limite]]
    if len(achados) > limite:
        textos.append(f'E mais {len(achados) - limite} conflito(s) de agenda: abra o ofício para ver todos.')
    return textos


def _data_final_do_oficio(form):
    """A data com que o ofício é finalizado e se ela foi posta pelo sistema.

    Vale a digitada em "Data do ofício" quando a pessoa a mudou; sem mudança
    (ou em branco), a de hoje — o dia em que o ofício sai.
    """
    from django.utils import timezone
    digitada = (form.data.get('data_criacao') or '').strip()
    if digitada and 'data_criacao' in form.changed_data:
        return form.cleaned_data['data_criacao'], False
    return timezone.localdate(), True


@acesso_ao_modulo
@require_http_methods(['GET', 'POST'])
def editar(request, pk=None):
    """Cadastro de ofício: as quatro etapas do Gerenciador de Viagens numa página.

    Dados e viajantes, roteiro e diárias (o editor da tela de roteiros),
    justificativa e documentos. "Salvar rascunho" grava e volta para a lista;
    "Finalizar Ofício" grava, confere as pendências e só finaliza sem elas.
    """
    from core.retorno import next_valido, voltar_para
    from .form_context import contexto_conferencia, contexto_dados_viajantes, contexto_justificativa
    from .justificativas_services import get_or_create_justificativa_oficio, oficio_exige_justificativa
    from .campos_modelo import aplicar, preencher_marcadores_do_oficio, valores_do_oficio
    from .presenters import artefatos_pdf_por_oficio, tipo_do_oficio
    from .protocolo_services import abrir_protocolo_do_oficio, mensagens_do_protocolo
    from .services import criar_oficio_rascunho, dados_eprotocolo
    exigir_operador(request)
    if pk is None:
        # O cadastro sempre edita um rascunho já numerado; sem ele, cria-se um.
        if request.method != 'POST':
            return redirect('viagens_oficios:lista')
        oficio = get_oficio_by_id(criar_oficio_rascunho().pk)
    else:
        oficio = get_oficio_by_id(pk)
    lista = voltar_para(request, _url_de_volta(oficio))
    finalizar = request.POST.get('acao') == 'finalizar'
    vincular = request.POST.get('acao') == 'vincular_roteiro'
    from viagens_cadastros.models import ConfiguracaoSistema
    unidade_emissora = ConfiguracaoSistema.para_usuario(request.user).unidade_id
    form = OficioForm(request.POST or None, instance=oficio, unidade_emissora=unidade_emissora)
    jform = JustificativaForm(
        request.POST or None, instance=get_or_create_justificativa_oficio(oficio), prefix='justificativa',
        obrigatoria=finalizar and oficio_exige_justificativa(oficio),
    )
    gravacao = None
    data_anterior = oficio.data_criacao
    data_automatica = False
    if request.method == 'POST':
        form_ok = form.is_valid()
        if form_ok and finalizar:
            # A data final do ofício vem ANTES da conferência: a digitada, ou a
            # de hoje se o campo não mudou. É com ela que se decide Autorização
            # ou Convalidação e se a justificativa é obrigatória.
            data_final, data_automatica = _data_final_do_oficio(form)
            form.instance.data_criacao = data_final
            jform._obrigatoria = oficio_exige_justificativa(form.instance)
        if form_ok and jform.is_valid():
            with transaction.atomic():
                oficio = form.save()
                reservar_numero_oficio(oficio, ano=oficio.data_criacao.year)
                if vincular:
                    # Escolher o roteiro na busca: grava o que já foi digitado e recarrega com ele.
                    vinculado = _vincular_roteiro(request, oficio)
                elif request.POST.get('vincular_roteiro') == '1' and request.POST.get('roteiro_trocado') == '1':
                    # Roteiro escolhido na tela: o editor mostra uma cópia dele, então só se vincula.
                    _vincular_roteiro(request, oficio)
                elif request.POST.get('vincular_roteiro') == '1' and not oficio.roteiro_id:
                    pass  # vínculo ligado sem roteiro escolhido: nada a gravar no roteiro
                else:
                    gravacao = _gravar_roteiro(request, oficio, finalizar=finalizar)
                atualizar_justificativa_oficio(oficio, jform, action='save_continue' if finalizar else 'save_draft')
                # Campos automáticos que ficaram marcados ({destino}...) viram valor agora que tudo está gravado.
                preencher_marcadores_do_oficio(oficio)
            # Fora da transação de propósito: abrir o protocolo é uma chamada
            # a outro sistema e não pode segurar a gravação do ofício — se
            # falhar, o ofício já está salvo e a tela explica o que houve.
            resultado_protocolo = abrir_protocolo_do_oficio(oficio)
            for nivel, texto in mensagens_do_protocolo(resultado_protocolo, finalizar=finalizar):
                messages.add_message(request, nivel, texto)
            if vincular:
                if vinculado is not None:
                    messages.success(request, f'Roteiro #{vinculado.pk} vinculado ao ofício. Rascunho salvo.')
                else:
                    messages.error(request, 'Escolha um roteiro existente para vincular.')
                return redirect(reverse('viagens_oficios:editar', args=[oficio.pk]) + '#roteiro')
            if oficio.protocolo:
                # Mesmo protocolo em outro ofício pode ser engano de digitação: avisa sem impedir.
                outros = (Oficio.objects.filter(protocolo=oficio.protocolo, cancelado=False)
                          .exclude(pk=oficio.pk).order_by('ano', 'numero'))
                if outros.exists():
                    from core.utils.masks import format_protocolo
                    nomes = ', '.join(o.numero_formatado for o in outros[:3])
                    messages.warning(request, f'O protocolo {format_protocolo(oficio.protocolo)} também está no ofício {nomes}.'
                                     if outros.count() == 1 else
                                     f'O protocolo {format_protocolo(oficio.protocolo)} também está nos ofícios {nomes}.')
            aviso_condutor = _aviso_de_condutor(oficio)
            if aviso_condutor:
                messages.warning(request, aviso_condutor)
            for texto in _avisos_de_conflito(oficio):
                messages.warning(request, texto)
            for nivel, texto in (gravacao.mensagens if gravacao else []):
                if texto.startswith('Diárias: R$'):
                    # O roteiro calcula por servidor; o aviso fala do ofício inteiro.
                    oficio.refresh_from_db()
                    diarias = _diarias_do_oficio(oficio)
                    if diarias:
                        pessoas = diarias['servidores']
                        texto = (f"Diárias: R$ {diarias['valor']} ({diarias['resumo']}) para "
                                 f"{pessoas} servidor{'es' if pessoas != 1 else ''}.")
                messages.add_message(request, nivel, texto)
            if gravacao is not None and not gravacao.ok:
                messages.error(request, 'Ofício salvo, mas o roteiro tem campos a corrigir.')
            elif finalizar:
                pendencias = validar_oficio_para_documento(oficio)['pendencias']
                if pendencias:
                    if data_automatica and oficio.data_criacao != data_anterior:
                        # Não finalizou: o rascunho fica com a data que tinha.
                        oficio.data_criacao = data_anterior
                        oficio.save(update_fields=['data_criacao', 'atualizado_em'])
                    for pendencia in pendencias:
                        messages.error(request, pendencia)
                    return redirect('viagens_oficios:editar', pk=oficio.pk)
                oficio.status = Oficio.STATUS_FINALIZADO
                oficio.save(update_fields=['status', 'atualizado_em'])
                messages.success(request, 'Ofício finalizado com sucesso.')
                if oficio.ano and oficio.ano != oficio.data_criacao.year:
                    # Ex.: rascunho de dezembro finalizado em janeiro.
                    messages.warning(request, (
                        f'O ofício {oficio.numero_formatado} tem número de {oficio.ano}, mas a data é de '
                        f'{oficio.data_criacao:%d/%m/%Y}. Confira se ele deve ser renumerado no ano da data.'))
                return redirect(lista)
            else:
                messages.success(request, 'Rascunho salvo.')
                return redirect(lista)
        else:
            messages.error(request, 'Não foi possível salvar o ofício. Revise os campos indicados.')
    oficio = get_oficio_by_id(oficio.pk)
    campos = valores_do_oficio(oficio)
    conferencia = contexto_conferencia(oficio, artefatos_pdf_por_oficio([oficio]).get(oficio.pk, {}))
    return render(request, 'pages/viagens_oficios/form.html', {
        'titulo': 'Cadastro de ofício',
        'oficio': oficio, 'form': form, 'jform': jform,
        'dados': contexto_dados_viajantes(form, oficio),
        'conflitos': _conflitos_da_tela(oficio, form),
        'conflitos_fixos': f'oficio={oficio.pk}',
        'rot': _contexto_roteiro(request, oficio, gravacao),
        'justificativa': contexto_justificativa(jform),
        'conferencia': conferencia,
        'tipo': tipo_do_oficio(oficio),
        'eprotocolo': dados_eprotocolo(oficio),
        'next': next_valido(request),
        'url_voltar': lista,
        'url_atual': request.get_full_path(),
        # O texto de cada modelo já com os campos automáticos do ofício.
        'modelos_texto': {
            'modelo_motivo': {pk: aplicar(texto, campos) for pk, texto in ModeloMotivoOficio.objects.values_list('pk', 'texto')},
            'justificativa-modelo': {pk: aplicar(texto, campos) for pk, texto in ModeloJustificativa.objects.values_list('pk', 'texto')},
        },
    })


@acesso_ao_modulo
@require_GET
def oficios_do_motorista(request, pk):
    """JSON do cartão do motorista: os ofícios ativos em que ele viaja, para
    preencher sozinho o N° do Ofício e o Protocolo de origem."""
    from django.http import JsonResponse
    from .services import oficios_do_motorista as buscar
    exigir_operador(request)
    oficio = get_oficio_by_id(pk)
    motorista = request.GET.get('motorista', '')
    if not motorista.isdigit():
        return JsonResponse({'oficios': []})
    return JsonResponse({'oficios': buscar(oficio, int(motorista))})


@acesso_ao_modulo
@require_GET
@xframe_options_sameorigin
def visualizar(request, pk, tipo):
    """O PDF do ofício ou da justificativa dentro do cartão da conferência."""
    if tipo not in ('oficio', 'justificativa'):
        raise Http404
    exigir_operador(request)
    oficio = get_oficio_by_id(pk)
    try:
        doc = gerar_documento(oficio, DocumentoFormato.PDF, DocumentoTipo(tipo))
    except (ValidationError, DocumentError) as exc:
        return HttpResponse('; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc), status=409, content_type='text/plain; charset=utf-8')
    return build_inline_pdf_response(request, content=doc.conteudo, tipo=doc.tipo, cache_hit=doc.cache_hit, x_document_sha256=doc.hash_sha256)


@acesso_ao_modulo
@require_GET
@xframe_options_sameorigin
def visualizar_termo(request, pk, servidor_id):
    """O termo de um servidor dentro do cartão da conferência."""
    from viagens_termos.services import gerar_termo_um
    exigir_operador(request)
    oficio = get_oficio_by_id(pk)
    if oficio.cancelado:
        raise Http404
    servidor = get_object_or_404(oficio.servidores_termo_autorizacao, pk=servidor_id)
    try:
        doc = gerar_termo_um(oficio, servidor, DocumentoFormato.PDF)
    except (ValidationError, DocumentError) as exc:
        return HttpResponse('; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc), status=409, content_type='text/plain; charset=utf-8')
    return build_inline_pdf_response(request, content=doc.conteudo, tipo=doc.tipo, cache_hit=doc.cache_hit, x_document_sha256=doc.hash_sha256)


@acesso_ao_modulo
@require_POST
def acao(request, pk, acao):
    from core.retorno import voltar_para
    from .services import OficioVinculadoError, desfazer_complementar_oficio, desfazer_retificacao_oficio
    exigir_operador(request)
    oficio = get_oficio_by_id(pk)
    # Da lista a ação volta para a lista como estava; do formulário, para o formulário.
    destino = voltar_para(request, reverse('viagens_oficios:editar', args=[pk]))
    if acao == 'cancelar':
        oficio.cancelar(request.POST.get('motivo', ''))
        messages.success(request, f'Ofício {oficio.numero_formatado} cancelado. O histórico foi mantido.')
    elif acao == 'reativar':
        oficio.reativar()
        messages.success(request, f'Ofício {oficio.numero_formatado} reativado.')
    elif acao == 'arquivar':
        oficio.status = Oficio.STATUS_ARQUIVADO
        oficio.save(update_fields=['status', 'atualizado_em'])
        messages.success(request, f'Ofício {oficio.numero_formatado} arquivado.')
    elif acao == 'retificar':
        # O mesmo item do menu liga e desliga a marca, como o estado de retificação da origem.
        if oficio.retificado_documento:
            desfazer_retificacao_oficio(oficio)
            messages.success(request, f'Ofício {oficio.numero_formatado} deixou de ser retificado.')
        else:
            retificar_oficio(oficio)
            messages.success(request, f'Ofício {oficio.numero_formatado} marcado como retificado.')
    elif acao == 'complementar':
        if oficio.complementar_documento:
            desfazer_complementar_oficio(oficio)
            messages.success(request, f'Ofício {oficio.numero_formatado} deixou de ser complementar.')
        else:
            marcar_oficio_complementar(oficio)
            messages.success(request, f'Ofício {oficio.numero_formatado} identificado como complementar.')
    elif acao == 'excluir':
        numero = oficio.numero_formatado
        volta = _url_de_volta(oficio)
        try:
            excluir_oficio(oficio)
        except OficioVinculadoError:
            messages.error(request, f'O ofício {numero} tem prestação de contas ou documentos vinculados e não pode ser excluído.')
            return redirect(destino)
        messages.success(request, f'Ofício {numero} excluído. O número volta para a sequência.')
        return redirect(voltar_para(request, volta))
    else:
        raise Http404
    return redirect(destino)


def resposta_documento(request, doc):
    if request.GET.get('inline') == '1' and doc.formato == DocumentoFormato.PDF:
        return build_inline_pdf_response(request, content=doc.conteudo, tipo=doc.tipo,
            cache_hit=doc.cache_hit, x_document_sha256=doc.hash_sha256)
    return build_download_response(content=doc.conteudo, tipo=doc.tipo, formato=doc.formato, cache_hit=doc.cache_hit)


@acesso_ao_modulo
@require_POST
def gerar(request, pk, tipo, formato):
    exigir_operador(request)
    if tipo not in ['oficio', 'justificativa'] or formato not in ['docx', 'pdf']:
        raise Http404
    oficio = get_oficio_by_id(pk)
    try:
        doc = gerar_documento(oficio, DocumentoFormato(formato), DocumentoTipo(tipo))
    except (ValidationError, DocumentError) as exc:
        messages.error(request, '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc))
        return redirect('viagens_oficios:editar', pk=pk)
    return resposta_documento(request, doc)


@acesso_ao_modulo
@require_GET
def documento(request, pk):
    """Endereço antigo da tela do documento: o editor agora é o visualizador
    no fim do formulário do ofício, no cartão do documento."""
    get_oficio_by_id(pk)
    return redirect(reverse('viagens_oficios:editar', args=[pk]) + '#documento-oficio')


@acesso_ao_modulo
@require_GET
@xframe_options_sameorigin
def documento_folha(request, pk):
    """O documento em si, no modo `editor`: o que o iframe da prévia mostra.

    A montagem mora no vínculo do editor, que é a mesma usada na resposta de
    gravação — a folha que abre e a que se atualiza saem do mesmo lugar.
    """
    from documentos.editor.vinculos import vinculo_do_tipo
    oficio = get_oficio_by_id(pk)
    html = vinculo_do_tipo(DocumentoTipo.OFICIO).folha(oficio, request.user)
    response = HttpResponse(html)
    response['Cache-Control'] = 'no-store'
    return response


@acesso_ao_modulo
@require_POST
def termos_todos_pdf(request, pk):
    """Todos os termos do ofício num PDF só — `baixar_termos_todos_pdf` da origem."""
    from viagens_termos.services import gerar_termo_lote
    from viagens_termos.views import resposta_pdf_consolidado
    exigir_operador(request)
    oficio = get_oficio_by_id(pk)
    if oficio.cancelado:
        messages.error(request, 'Reative o ofício antes de gerar termos.')
        return redirect('viagens_oficios:editar', pk=pk)
    try:
        documentos = gerar_termo_lote(oficio, DocumentoFormato.PDF)
    except (ValidationError, DocumentError) as exc:
        messages.error(request, '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc))
        return redirect('viagens_oficios:editar', pk=pk)
    if not documentos:
        messages.error(request, 'Nenhum servidor selecionado para termo neste ofício.')
        return redirect('viagens_oficios:editar', pk=pk)
    return resposta_pdf_consolidado(documentos, f"{oficio.numero_formatado.replace('/', '-')}-termos.pdf")


@acesso_ao_modulo
@require_POST
def termos(request, pk, formato, servidor_id=None):
    from viagens_termos.services import gerar_termo_um, gerar_termo_lote
    exigir_operador(request)
    if formato not in ['docx', 'pdf']:
        raise Http404
    oficio = get_oficio_by_id(pk)
    if oficio.cancelado:
        messages.error(request, 'Reative o ofício antes de gerar termos.')
        return redirect('viagens_oficios:editar', pk=pk)
    fmt = DocumentoFormato(formato)
    try:
        if servidor_id:
            servidor = get_object_or_404(oficio.servidores_termo_autorizacao, pk=servidor_id)
            return resposta_documento(request, gerar_termo_um(oficio, servidor, fmt))
        return resposta_lote(gerar_termo_lote(oficio, fmt))
    except (ValidationError, DocumentError) as exc:
        messages.error(request, '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc))
        return redirect('viagens_oficios:editar', pk=pk)


def resposta_lote(documentos, nome='termos.zip'):
    buffer = io.BytesIO()
    with ZipFile(buffer, 'w') as zipfile:
        for doc in documentos:
            zipfile.writestr(doc.nome_arquivo, doc.conteudo)
    response = HttpResponse(buffer.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="{nome}"'
    response['Cache-Control'] = 'no-store'
    return response


@acesso_ao_modulo
@require_http_methods(['GET', 'POST'])
def numeracao(request):
    if not eh_gestor_viagens(request.user):
        raise PermissionDenied
    from django.utils import timezone
    ano = request.POST.get('ano') or request.GET.get('ano') or timezone.localdate().year
    cfg = ConfiguracaoNumeracaoOficio.objects.filter(ano=int(ano)).first() if str(ano).isdigit() else None
    form = NumeracaoForm(request.POST or None, instance=cfg, initial={'ano': ano})
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            bloquear_escopo_numeracao(namespace=NAMESPACE_OFICIO, ano=form.cleaned_data['ano'], modelo=Oficio)
            form.save()
        messages.success(request, 'Piso de numeração salvo.')
        return redirect('viagens_oficios:numeracao')
    return render(request, 'pages/viagens_oficios/form_simples.html', {
        'titulo': 'Numeração de ofícios', 'form': form, 'campos': campos_v32(form),
        'ajuda': 'A sequência é calculada pelo maior número ocupado e pelas lacunas liberadas. O piso não renumera ofícios existentes.',
        'url_voltar': reverse('viagens_oficios:lista'),
    })
