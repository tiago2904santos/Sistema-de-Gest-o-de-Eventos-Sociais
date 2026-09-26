import logging
from datetime import timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.base import ContentFile
from django.core.paginator import Paginator
from django.db import DatabaseError, transaction
from django.db.models import Count, F, Q
from django.http import FileResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import (
    AnexoForm,
    DespachoForm,
    FiltroSolicitacoesForm,
    SolicitacaoForm,
    validar_arquivo_anexo,
)
from .models import (
    AcaoHistorico,
    AnexoSolicitacao,
    DecisaoDG,
    HistoricoSolicitacao,
    SolicitacaoEvento,
    StatusSolicitacao,
    TipoOperacao,
)
from core import preencher_por_email
from integracoes.eprotocolo import andamento as andamento_eprotocolo
from integracoes.eprotocolo.andamento import formatar_numero
from core.listagens import trilha_de_situacoes

from .presenters import linha_da_lista
from . import integracao_viagens, lembretes, permissions, preenchimento, services

ITENS_POR_PAGINA = 15

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Helpers de contexto para os components do design system
# ---------------------------------------------------------------------------

def _opcoes(iteravel):
    return [{"valor": str(getattr(item, "pk", item)), "rotulo": str(item)} for item in iteravel]


def _opcoes_municipios(iteravel):
    return [
        {
            "valor": str(item.pk),
            "rotulo": str(item),
            "estado": str(item.estado_id),
        }
        for item in iteravel
    ]


def _opcoes_choices(choices):
    return [{"valor": str(valor), "rotulo": str(rotulo)} for valor, rotulo in choices]


def _valores(form):
    """Valores atuais (POST ou instância) como strings para os templates."""
    valores = {}
    for nome in form.fields:
        valor = form[nome].value()
        if isinstance(valor, bool):
            valor = "1" if valor else "0"
        valores[nome] = "" if valor is None else str(valor)
    return valores


def _marcados(form, nome):
    """IDs selecionados de um campo múltiplo, como strings."""
    valor = form[nome].value() or []
    return [str(getattr(item, "pk", item)) for item in valor]


CAMPOS_FORMULARIO = [
    "data_solicitacao", "data_inicio_evento", "data_fim_evento", "tipo_evento",
    "municipio", "local_evento", "protocolo", "solicitante_nome", "solicitante_cargo_unidade",
    "contato", "orgao_responsavel", "unidade_movel", "unidade_movel_designada",
    "descricao_complementar", "quantidade_servidores",
    "tipo_operacao", "quantidade_cin", "motorista", "decisao_dg", "observacoes_dg",
]

CAMPOS_FK = {
    "tipo_evento", "municipio", "orgao_responsavel", "motorista",
    "unidade_movel_designada",
}


def _valor_da_instancia(solicitacao, nome):
    if nome in CAMPOS_FK:
        valor = getattr(solicitacao, f"{nome}_id")
    else:
        valor = getattr(solicitacao, nome)
    if isinstance(valor, bool):
        return "1" if valor else "0"
    return "" if valor is None else str(valor)


def _breadcrumb(titulo):
    """Trilha das telas de solicitação, sempre passando pela listagem."""
    return [
        {"label": "Solicitações", "url": reverse("solicitacoes:lista")},
        {"label": titulo},
    ]


def _contexto_formulario(request, form, solicitacao=None, reabrindo=False):
    if solicitacao:
        acoes = permissions.acoes_permitidas(request.user, solicitacao)
        if reabrindo:
            # No "Editar" a tela é a da etapa 1; o único caminho é reenviar.
            acoes = {**acoes, "editar_dados": True, "enviar": False,
                     "despachar": False, "concluir": False,
                     "aguarda_atendimento": False, "cancelar": False}
    else:
        acoes = {"editar_dados": True, "enviar": True, "despachar": False}

    valores = _valores(form)
    # Campos fora do formulário (seções desabilitadas) exibem o valor salvo.
    if solicitacao:
        for nome in CAMPOS_FORMULARIO:
            if nome not in valores:
                valores[nome] = _valor_da_instancia(solicitacao, nome)
    if "estado" not in valores:
        valores["estado"] = (
            str(solicitacao.municipio.estado_id)
            if solicitacao and solicitacao.municipio_id
            else ""
        )

    def opcoes_de(nome, fallback_relacao=None):
        if nome in form.fields:
            return _opcoes(form.fields[nome].queryset)
        if solicitacao and fallback_relacao:
            return _opcoes(fallback_relacao)
        if solicitacao and getattr(solicitacao, nome, None):
            return _opcoes([getattr(solicitacao, nome)])
        return []

    def marcados_de(nome, relacao):
        if nome in form.fields:
            return _marcados(form, nome)
        if solicitacao:
            return [str(item.pk) for item in relacao]
        return []

    servicos_salvos = list(solicitacao.servicos.all()) if solicitacao else []
    equipes_salvas = list(solicitacao.equipes.all()) if solicitacao else []
    estados_salvos = (
        [solicitacao.municipio.estado]
        if solicitacao and solicitacao.municipio_id
        else []
    )
    equipes_disponiveis = opcoes_de("equipes", equipes_salvas)
    equipes_marcadas = marcados_de("equipes", equipes_salvas)
    quantidades_salvas = {
        str(item.equipe_id): item.quantidade_servidores
        for item in solicitacao.itens_equipe.all()
    } if solicitacao else {}
    equipes_planejamento = []
    for equipe in equipes_disponiveis:
        nome_quantidade = f"quantidade_equipe_{equipe['valor']}"
        if form.is_bound:
            quantidade = form.data.get(nome_quantidade, "")
        else:
            quantidade = quantidades_salvas.get(equipe["valor"], "")
        equipes_planejamento.append(
            {
                **equipe,
                "selecionada": equipe["valor"] in equipes_marcadas,
                "nome_quantidade": nome_quantidade,
                # A DG ajusta a mesma equipe num campo próprio, no despacho.
                "nome_quantidade_dg": f"quantidade_dg_{equipe['valor']}",
                "quantidade": "" if quantidade is None else str(quantidade),
            }
        )

    servicos_marcados = marcados_de("servicos", servicos_salvos)
    # Os serviços viram cartões de escolha: cada um já sabe se está marcado.
    servicos_cartoes = [
        {**servico, "marcado": servico["valor"] in servicos_marcados}
        for servico in opcoes_de("servicos", servicos_salvos)
    ]

    return {
        "form": form,
        "solicitacao": solicitacao,
        "acoes": acoes,
        # Cabeçalho da tela de edição: só o título e o selo da situação.
        "selo": solicitacao.get_status_display() if solicitacao else "Rascunho",
        "selo_tom": solicitacao.status.lower() if solicitacao else "rascunho",
        "valores": valores,
        "erros": form.errors,
        "erro_periodo": form.errors.get("data_inicio_evento")
        or form.errors.get("data_fim_evento"),
        "tipos_evento": opcoes_de("tipo_evento"),
        "estados": opcoes_de("estado", estados_salvos),
        "municipios": _opcoes_municipios(form.fields["municipio"].queryset)
        if "municipio" in form.fields
        else _opcoes_municipios([solicitacao.municipio])
        if solicitacao and solicitacao.municipio_id
        else [],
        "orgaos": opcoes_de("orgao_responsavel"),
        "servicos": servicos_cartoes,
        "equipes": equipes_disponiveis,
        "equipes_planejamento": equipes_planejamento,
        "motoristas": opcoes_de("motorista"),
        "unidades_moveis": opcoes_de("unidade_movel_designada"),
        "tipos_operacao": _opcoes_choices(TipoOperacao.choices),
        "servicos_marcados": servicos_marcados,
        "equipes_marcadas": equipes_marcadas,
        "timeline": services.montar_timeline(solicitacao),
        "dados_desabilitado": bool(solicitacao) and not acoes["editar_dados"],
        # Formulário único: o planejamento acompanha a edição dos dados.
        "planejamento_desabilitado": bool(solicitacao) and not acoes["editar_dados"],
        "mostrar_enviar": acoes["enviar"] if solicitacao else True,
        # Só aparece para quem pode decidir agora: fora disso a decisão da DG
        # já está no resumo, e um formulário desabilitado só confunde.
        "mostrar_despacho_dg": bool(solicitacao) and acoes.get("despachar", False),
        "anexos": list(solicitacao.anexos.select_related("enviado_por"))
        if solicitacao
        else [],
        "pode_gerenciar_anexos": acoes.get("gerenciar_anexos", False)
        if solicitacao
        else False,
    }


def _obter_visivel(request, pk):
    solicitacao = get_object_or_404(
        SolicitacaoEvento.objects.select_related(
            "municipio__estado", "regiao", "tipo_evento", "orgao_responsavel", "motorista",
            "criado_por", "decidido_por",
        ).prefetch_related(
            "servicos", "equipes", "itens_equipe__equipe", "historico__usuario"
        ),
        pk=pk,
    )
    if not permissions.pode_ver(request.user, solicitacao):
        raise PermissionDenied
    return solicitacao


# ---------------------------------------------------------------------------
# Criação e edição
# ---------------------------------------------------------------------------

def _criar_anexos_enviados(solicitacao, arquivos, usuario):
    for arquivo in arquivos:
        AnexoSolicitacao.objects.create(
            solicitacao=solicitacao,
            arquivo=arquivo,
            nome_original=arquivo.name,
            tamanho=arquivo.size,
            enviado_por=usuario,
        )


def _anexar_email_de_origem(solicitacao, origem, usuario):
    """Anexa o e-mail de onde a solicitação saiu e os anexos dele (o ofício em PDF…).

    Cada arquivo passa pela mesma validação do anexo comum (tipo real e
    tamanho); o que não passar fica de fora e volta como aviso para a tela.
    O texto colado entra como .txt: é o único registro do pedido.
    """
    arquivos = [(origem.nome, origem.dados, "O e-mail")]
    mensagem = None if origem.colado else origem.mensagem()
    if mensagem is not None:
        arquivos += [(nome, dados, f"O anexo {nome} do e-mail") for nome, dados in mensagem.anexos]
    avisos = []
    for nome, dados, descricao in arquivos:
        arquivo = ContentFile(dados, name=nome)
        erro = validar_arquivo_anexo(arquivo)
        if erro:
            avisos.append(f"{descricao} não foi anexado — {erro}")
            continue
        AnexoSolicitacao.objects.create(
            solicitacao=solicitacao,
            arquivo=arquivo,
            nome_original=nome[:255],
            tamanho=len(dados),
            enviado_por=usuario,
        )
    return avisos


def _duplicados_do_email(usuario):
    """Solicitações que este usuário vê e que já saíram do mesmo e-mail."""

    def buscar(texto_origem):
        historicos = (
            HistoricoSolicitacao.objects.filter(
                acao=AcaoHistorico.CRIACAO, observacao=texto_origem
            )
            .select_related("solicitacao")
            .order_by("-criado_em")[:5]
        )
        return [
            {
                "titulo": f"Solicitação #{h.solicitacao.pk}",
                "url": reverse("solicitacoes:editar", args=[h.solicitacao.pk]),
            }
            for h in historicos
            if permissions.pode_ver(usuario, h.solicitacao)
        ]

    return buscar


@login_required
@require_POST
def ler_email(request):
    """Lê o e-mail do pedido (arquivo ou texto colado) para a tela "Nova solicitação".

    Não grava nada: devolve as sugestões em JSON (`core.preencher_por_email`)
    e guarda o original até a solicitação ser salva.
    """
    return preencher_por_email.responder_leitura(
        request,
        modulo="solicitacoes",
        sugerir=preenchimento.sugestoes,
        formulario=SolicitacaoForm,
        anexa_original=True,
        duplicados=_duplicados_do_email(request.user),
    )


@login_required
def nova_solicitacao(request):
    """Tela "Nova Solicitação de Evento Social" com persistência real."""
    email_origem = None
    if request.method == "POST":
        acao = request.POST.get("acao", "rascunho")
        form = SolicitacaoForm(request.POST, enviar=(acao == "enviar"))
        arquivos = request.FILES.getlist("anexos")
        erros_anexos = [
            erro
            for erro in (validar_arquivo_anexo(arquivo) for arquivo in arquivos)
            if erro
        ]
        for erro in erros_anexos:
            messages.error(request, erro)
        # O e-mail lido em "Preencher com um e-mail", se a tela veio dele.
        origem = preencher_por_email.origem_do_pedido(request, "solicitacoes")
        if form.is_valid() and not erros_anexos:
            avisos_origem = []
            with transaction.atomic():
                solicitacao = form.save(criado_por=request.user)
                services.registrar_historico(
                    solicitacao, request.user, AcaoHistorico.CRIACAO,
                    status_novo=solicitacao.status,
                    observacao=preencher_por_email.texto_da_origem(origem) if origem else "",
                )
                _criar_anexos_enviados(solicitacao, arquivos, request.user)
                if origem:
                    avisos_origem = _anexar_email_de_origem(solicitacao, origem, request.user)
                if acao == "enviar":
                    services.enviar(solicitacao, request.user)
            preencher_por_email.concluir_origem(request, origem)
            for aviso in avisos_origem:
                messages.warning(request, aviso)
            if acao == "enviar":
                messages.success(request, f"Solicitação #{solicitacao.pk} enviada com sucesso.")
            else:
                messages.success(request, f"Rascunho #{solicitacao.pk} salvo com sucesso.")
            return redirect("solicitacoes:editar", pk=solicitacao.pk)
        else:
            messages.error(request, "Corrija os campos destacados para continuar.")
            email_origem = preencher_por_email.origem_pendente(request, "solicitacoes")
    else:
        form = SolicitacaoForm()
    contexto = _contexto_formulario(request, form)
    contexto["email_origem"] = email_origem
    contexto["titulo_pagina"] = "Nova Solicitação de Evento Social"
    contexto["breadcrumb"] = _breadcrumb("Nova solicitação")
    return render(request, "pages/solicitacoes/form.html", contexto)


@login_required
def editar_solicitacao(request, pk):
    """Tela única do registro: dados, workflow, anexos e histórico.

    Quem pode ver abre a tela; quem não pode editar recebe os campos
    desabilitados, mas continua com as ações de workflow a que tem direito
    (despacho da DG, encerramento, anexos). Só o POST exige alçada de edição.
    """
    solicitacao = _obter_visivel(request, pk)
    # "Editar" de um pedido já enviado (`?reabrir=1`): a tela da etapa 1 abre
    # editável, e salvar devolve o pedido para novo despacho da DG.
    parametros = request.POST if request.method == "POST" else request.GET
    reabrindo = parametros.get("reabrir") == "1" and permissions.pode_reabrir(
        request.user, solicitacao
    )
    pode_editar = reabrindo or permissions.pode_editar_dados(request.user, solicitacao)

    if request.method == "POST":
        if not pode_editar:
            raise PermissionDenied
        acao = "reenviar" if reabrindo else request.POST.get("acao", "rascunho")
        # Antes do form: a validação já escreve os valores novos na instância.
        antes = services.fotografia(solicitacao)
        if acao == "enviar" and not permissions.pode_enviar(request.user, solicitacao):
            raise PermissionDenied
        form = SolicitacaoForm(
            request.POST, instance=solicitacao, enviar=acao in ("enviar", "reenviar")
        )
        if form.is_valid():
            try:
                with transaction.atomic():
                    solicitacao = form.save()
                    alteracoes = services.diferencas(
                        antes, services.fotografia(solicitacao)
                    )
                    if acao == "reenviar":
                        if alteracoes:
                            services.reenviar_apos_edicao(
                                solicitacao, request.user, alteracoes
                            )
                    else:
                        if alteracoes:
                            services.registrar_historico(
                                solicitacao,
                                request.user,
                                AcaoHistorico.ATUALIZACAO,
                                alteracoes=alteracoes,
                            )
                        if acao == "enviar":
                            services.enviar(solicitacao, request.user)
            except ValidationError as erro:
                for mensagem_erro in erro.messages:
                    messages.error(request, mensagem_erro)
            else:
                if acao == "reenviar" and not alteracoes:
                    messages.info(request, "Nada foi alterado: a solicitação continua como estava.")
                elif acao == "reenviar":
                    messages.success(
                        request,
                        f"Solicitação #{solicitacao.pk} alterada e reenviada para o despacho da DG.",
                    )
                elif acao == "enviar":
                    messages.success(
                        request, f"Solicitação #{solicitacao.pk} enviada com sucesso."
                    )
                else:
                    messages.success(request, f"Solicitação #{solicitacao.pk} atualizada.")
                return redirect("solicitacoes:editar", pk=solicitacao.pk)
        else:
            messages.error(request, "Corrija os campos destacados para continuar.")
    else:
        form = SolicitacaoForm(instance=solicitacao)

    devolucao = (
        solicitacao.historico.filter(acao=AcaoHistorico.DEVOLUCAO).last()
        if solicitacao.status == StatusSolicitacao.DEVOLVIDA
        else None
    )
    # O despacho pendente sai da sessão ao ser lido: uma leitura só.
    pendente = _despacho_pendente(request, solicitacao)
    contexto = _contexto_formulario(request, form, solicitacao, reabrindo=reabrindo)
    contexto.update(
        {
            "titulo_pagina": f"Solicitação #{solicitacao.pk}",
            "reabrindo": reabrindo,
            "breadcrumb": _breadcrumb(f"Solicitação #{solicitacao.pk}"),
            "somente_leitura": not pode_editar,
            "historico": solicitacao.historico.all(),
            "itens_equipe": list(solicitacao.itens_equipe.select_related("equipe")),
            "motivo_devolucao": devolucao,
            "despacho_pendente": pendente,
            "decisoes_dg": _decisoes_dg(pendente),
            **_contexto_despacho(request.user, solicitacao),
            "cartao_viagem": None if reabrindo else _cartao_viagem(request.user, solicitacao),
            "responsaveis": _opcoes_responsaveis(solicitacao)
            if permissions.pode_transferir(request.user, solicitacao) and not reabrindo
            else [],
            "andamento_protocolo": andamento_eprotocolo.andamento_guardado(
                request, "solicitacoes", solicitacao.pk
            ),
        }
    )
    return render(request, "pages/solicitacoes/form.html", contexto)


def _contexto_despacho(user, solicitacao):
    """Posição na fila ("3 de 7") e os textos prontos, só para quem despacha."""
    if not permissions.pode_despachar(user, solicitacao):
        return {}
    from cadastros.models import TextoDespacho

    fila = _fila_de_despacho(user)
    posicao = fila.index(solicitacao.pk) + 1 if solicitacao.pk in fila else None
    return {
        "fila_despacho": {"posicao": posicao, "total": len(fila)},
        "textos_despacho": list(TextoDespacho.objects.filter(ativo=True)),
    }


def _cartao_viagem(user, solicitacao):
    """O cartão "Viagem" da solicitação: o link, a situação e o que falta.

    Só aparece depois do deferimento (ou quando já existe viagem): antes
    disso não há logística a acompanhar.
    """
    from viagens_cadastros.permissions import pode_acessar

    viagem = integracao_viagens.viagem_da_solicitacao(solicitacao)
    if viagem is None and solicitacao.decisao_dg != DecisaoDG.ATENDER:
        return None
    if viagem is not None:
        return {
            "viagem": viagem,
            # Quem não tem o módulo vê a situação, mas não o link.
            "url": reverse("viagens_viagem:painel", args=[viagem.pk])
            if pode_acessar(user)
            else "",
            "falta": integracao_viagens.o_que_falta(viagem),
        }
    cabe, motivo = integracao_viagens.pode_gerar(solicitacao)
    return {
        "viagem": None,
        "pode_gerar": cabe and permissions.pode_gerar_viagem(user),
        "motivo": motivo,
    }


# ---------------------------------------------------------------------------
# Listagem e detalhe
# ---------------------------------------------------------------------------

# Filas = atalhos de status no mesmo escopo da tabela (todo mundo), exceto
# as pessoais, marcadas com "apenas_do_usuario" e rotuladas como "minhas".
FILAS = {
    "despacho": {
        "rotulo": "Aguardando despacho",
        "status": [StatusSolicitacao.AGUARDANDO_DESPACHO],
    },
    "devolvidas": {
        "rotulo": "Devolvidas para ajuste",
        "status": [StatusSolicitacao.DEVOLVIDA],
    },
    "andamento": {
        "rotulo": "Deferidas",
        "status": [StatusSolicitacao.DEFERIDA_EM_ANDAMENTO],
    },
    # Deferida com o evento já realizado: falta o autor marcar "Atendida".
    "confirmar": {
        "rotulo": "Confirmar atendimento",
        "status": [StatusSolicitacao.DEFERIDA_EM_ANDAMENTO],
        "evento_encerrado": True,
    },
    "canceladas": {
        "rotulo": "Canceladas",
        "status": [StatusSolicitacao.CANCELADA],
    },
    "rascunhos": {
        "rotulo": "Meus rascunhos",
        "status": [StatusSolicitacao.RASCUNHO],
        "apenas_do_usuario": True,
    },
    "minhas": {
        "rotulo": "Minhas",
        "apenas_do_usuario": True,
    },
    # Destinos dos cartões do Dashboard: fora da trilha, mas com o mesmo
    # recorte do número clicado. Aparecem como filtro ativo acima da lista.
    "deferidas_ano": {
        "rotulo": "Deferidas no ano",
        "status": [
            StatusSolicitacao.DEFERIDA_EM_ANDAMENTO,
            StatusSolicitacao.ATENDIDA,
        ],
        "ano_corrente": True,
        "oculta": True,
    },
    "proximos": {
        "rotulo": "Eventos nos próximos 30 dias",
        "proximos_dias": 30,
        "oculta": True,
    },
}


def _condicao_da_fila(config, user):
    """O recorte de uma fila como Q, o mesmo na contagem e na lista."""
    condicao = Q()
    if config.get("status"):
        condicao &= Q(status__in=config["status"])
    if config.get("apenas_do_usuario"):
        condicao &= Q(criado_por=user)
    if config.get("evento_encerrado"):
        condicao &= lembretes.condicao_evento_encerrado()
    if config.get("ano_corrente"):
        condicao &= Q(data_solicitacao__year=timezone.localdate().year)
    if config.get("proximos_dias"):
        hoje = timezone.localdate()
        condicao &= Q(
            data_inicio_evento__gte=hoje,
            data_inicio_evento__lte=hoje + timedelta(days=config["proximos_dias"]),
        ) & ~Q(
            status__in=[StatusSolicitacao.CANCELADA, StatusSolicitacao.NAO_ATENDIDA]
        )
    return condicao


def _filas_do_usuario(user, queryset):
    """Atalhos de fila com contagem, conforme o perfil do usuário."""
    # A fila de despacho é da DG; as pessoais valem para todos.
    filas = []
    if permissions.eh_gestor_dg(user):
        filas.append("despacho")
    filas.extend(
        ["devolvidas", "andamento", "confirmar", "canceladas", "rascunhos", "minhas"]
    )
    agregacoes = {
        chave: Count("pk", filter=_condicao_da_fila(FILAS[chave], user))
        for chave in filas
    }
    totais = queryset.aggregate(**agregacoes)
    resultado = []
    for chave in filas:
        config = FILAS[chave]
        resultado.append(
            {
                "chave": chave,
                "rotulo": config["rotulo"],
                "total": totais[chave],
            }
        )
    return resultado


# A ordenação não tem mais controles na tela (a lista segue a composição de
# Viagens), mas `?ordem=` continua valendo para os links antigos e para a
# exportação: rótulo da URL -> campos do banco.
ORDENACOES = {
    "numero": ["pk"],
    "municipio": ["municipio__nome", "-data_solicitacao"],
    "tipo": ["tipo_evento__nome", "-data_solicitacao"],
    "periodo": ["data_inicio_evento", "-pk"],
    "solicitante": ["solicitante_nome", "-data_solicitacao"],
    "data": ["data_solicitacao", "-pk"],
    "status": ["status", "-data_solicitacao"],
}
ORDENACAO_PADRAO = "-data"
# Na fila de despacho o que vence primeiro vem primeiro: o evento mais
# próximo no topo, os sem data no fim.
ORDEM_DESPACHO = [F("data_inicio_evento").asc(nulls_last=True), "pk"]


def _fila_de_despacho(user):
    """As solicitações aguardando despacho, na ordem de trabalho da DG."""
    return list(
        permissions.queryset_visivel(
            user,
            SolicitacaoEvento.objects.filter(
                status=StatusSolicitacao.AGUARDANDO_DESPACHO
            ),
        )
        .order_by(*ORDEM_DESPACHO)
        .values_list("pk", flat=True)
    )


def _ordenacao(request):
    """Lê `ordem` da URL (com "-" para decrescente) e devolve (chave, campos)."""
    pedido = request.GET.get("ordem") or ORDENACAO_PADRAO
    decrescente = pedido.startswith("-")
    chave = pedido.lstrip("-")
    if chave not in ORDENACOES:
        pedido, decrescente, chave = ORDENACAO_PADRAO, True, "data"
    campos = ORDENACOES[chave]
    if decrescente:
        campos = [
            campo[1:] if campo.startswith("-") else f"-{campo}" for campo in campos
        ]
    return pedido, campos


def _queryset_filtrado(request):
    """Solicitações visíveis com fila e filtros da listagem aplicados."""
    filtros = FiltroSolicitacoesForm(request.GET or None)
    base = permissions.queryset_visivel(
        request.user,
        SolicitacaoEvento.objects.select_related(
            "municipio", "tipo_evento", "regiao", "criado_por",
            # A linha mostra qual unidade móvel vai ao evento.
            "unidade_movel_designada",
        ),
    )
    _pedido, campos_ordem = _ordenacao(request)
    queryset = base.order_by(*campos_ordem)

    fila = request.GET.get("fila", "")
    if fila in FILAS:
        queryset = queryset.filter(_condicao_da_fila(FILAS[fila], request.user))
    else:
        fila = ""
    if fila == "despacho" and not request.GET.get("ordem"):
        queryset = queryset.order_by(*ORDEM_DESPACHO)

    if filtros.is_valid():
        dados = filtros.cleaned_data
        if dados.get("q"):
            termo = dados["q"]
            condicao = (
                Q(solicitante_nome__icontains=termo)
                | Q(local_evento__icontains=termo)
                | Q(municipio__nome__icontains=termo)
                | Q(protocolo__icontains=termo)
                | Q(tipo_evento__nome__icontains=termo)
                | Q(orgao_responsavel__nome__icontains=termo)
            )
            # "#123" ou "123": o número da solicitação.
            numero_solicitacao = termo.strip().lstrip("#").strip()
            if numero_solicitacao.isdigit() and len(numero_solicitacao) <= 9:
                condicao |= Q(pk=int(numero_solicitacao))
            # O número digitado sem pontos também acha o protocolo.
            numero = formatar_numero(termo)
            if numero:
                condicao |= Q(protocolo=numero)
            queryset = queryset.filter(condicao)
        if dados.get("status"):
            queryset = queryset.filter(status=dados["status"])
        if dados.get("municipio"):
            queryset = queryset.filter(municipio=dados["municipio"])
        if dados.get("tipo_evento"):
            queryset = queryset.filter(tipo_evento=dados["tipo_evento"])
        if dados.get("inicio"):
            queryset = queryset.filter(data_inicio_evento__gte=dados["inicio"])
        if dados.get("fim"):
            queryset = queryset.filter(data_inicio_evento__lte=dados["fim"])

    return queryset, base, filtros, fila


CAMPOS_FILTRO = ["q", "status", "municipio", "tipo_evento", "inicio", "fim"]


# Ícone de cada fila na trilha lateral das situações.
ICONES_FILA = {
    "despacho": "gavel",
    "devolvidas": "undo",
    "andamento": "check-circle",
    "confirmar": "check",
    "canceladas": "ban",
    "rascunhos": "pencil",
    "minhas": "user",
}


def _filtros_ativos(request, filtros, fila, filas_da_trilha):
    """Os filtros que a trilha não mostra, cada um com o "x" que o remove.

    Quem chega por um cartão do Dashboard (ou por um link salvo) traz status,
    período ou uma fila fora da trilha: sem isto o filtro ficava invisível e
    a lista parecia vazia ou errada.
    """
    base = request.GET.copy()
    base.pop("pagina", None)

    def sem(*nomes):
        destino = base.copy()
        for nome in nomes:
            destino.pop(nome, None)
        consulta = destino.urlencode()
        return f"?{consulta}" if consulta else "?"

    ativos = []
    if fila and fila not in filas_da_trilha:
        ativos.append({"rotulo": FILAS[fila]["rotulo"], "url_remover": sem("fila")})
    dados = filtros.cleaned_data if filtros.is_valid() else {}
    if dados.get("status"):
        ativos.append({
            "rotulo": f"Situação: {StatusSolicitacao(dados['status']).label}",
            "url_remover": sem("status"),
        })
    if dados.get("municipio"):
        ativos.append({
            "rotulo": f"Município: {dados['municipio']}",
            "url_remover": sem("municipio"),
        })
    if dados.get("tipo_evento"):
        ativos.append({
            "rotulo": f"Tipo: {dados['tipo_evento']}",
            "url_remover": sem("tipo_evento"),
        })
    if dados.get("inicio"):
        ativos.append({
            "rotulo": f"Eventos a partir de {dados['inicio']:%d/%m/%Y}",
            "url_remover": sem("inicio"),
        })
    if dados.get("fim"):
        ativos.append({
            "rotulo": f"Eventos até {dados['fim']:%d/%m/%Y}",
            "url_remover": sem("fim"),
        })
    return ativos


# Filtros que a busca carrega escondidos: buscar não pode descartar o
# período ou a situação que vieram do Dashboard.
CAMPOS_OCULTOS_NA_BUSCA = ["status", "municipio", "tipo_evento", "inicio", "fim", "ordem"]

# Os campos do painel "Filtros"; o resto da querystring vai escondido nele.
CAMPOS_DO_PAINEL = ["inicio", "fim", "municipio", "tipo_evento"]


def _painel_filtros(request, filtros):
    """O painel recolhível de período, município e tipo, pronto para a tela."""
    valores = {nome: request.GET.get(nome, "") for nome in CAMPOS_DO_PAINEL}
    total = sum(1 for valor in valores.values() if valor)
    ocultos = [
        {"nome": nome, "valor": valor}
        for nome, valor in request.GET.items()
        if nome not in CAMPOS_DO_PAINEL and nome != "pagina" and valor
    ]
    limpar = request.GET.copy()
    for nome in [*CAMPOS_DO_PAINEL, "pagina"]:
        limpar.pop(nome, None)
    consulta = limpar.urlencode()
    return {
        "valores": valores,
        "total": total,
        "aberto": bool(total),
        "ocultos": ocultos,
        "url_limpar": f"?{consulta}" if consulta else "?",
        "municipios": _opcoes(filtros.fields["municipio"].queryset),
        "tipos": _opcoes(filtros.fields["tipo_evento"].queryset.order_by("nome")),
    }


@login_required
def lista_solicitacoes(request):
    queryset, base, filtros, fila = _queryset_filtrado(request)

    paginador = Paginator(queryset, ITENS_POR_PAGINA)
    pagina = paginador.get_page(request.GET.get("pagina"))
    # Números de página com reticências ("1 2 3 … 7") para pular direto.
    paginas_visiveis = list(
        paginador.get_elided_page_range(pagina.number, on_each_side=1, on_ends=1)
    )

    parametros = request.GET.copy()
    parametros.pop("pagina", None)
    total_geral = base.count()
    filas = _filas_do_usuario(request.user, base)
    filtros_ativos = _filtros_ativos(
        request, filtros, fila, {item["chave"] for item in filas}
    )

    return render(
        request,
        "pages/solicitacoes/lista.html",
        {
            "pagina": pagina,
            "q": request.GET.get("q", ""),
            "querystring": parametros.urlencode(),
            # Trocar de situação na trilha substitui o status que veio do
            # Dashboard, em vez de somar os dois e esvaziar a lista.
            "situacoes": trilha_de_situacoes(
                request, filas, total_geral, ICONES_FILA, descartar=("status",)
            ),
            # Item aceso: a fila escolhida; "Todas" só quando nada filtra a
            # situação — com status na URL, "Todas" mentiria.
            "situacao_ativa": fila
            or ("" if request.GET.get("status") else "todas"),
            "fila_ativa": fila,
            "filtros_ativos": filtros_ativos,
            "campos_ocultos_busca": [
                {"nome": nome, "valor": request.GET[nome]}
                for nome in CAMPOS_OCULTOS_NA_BUSCA
                if request.GET.get(nome)
            ],
            "tem_filtros": bool(fila) or any(
                request.GET.get(nome) for nome in CAMPOS_FILTRO
            ),
            "paginas_visiveis": paginas_visiveis,
            "elipse": Paginator.ELLIPSIS,
            "painel": _painel_filtros(request, filtros),
            "url_exportar": reverse("solicitacoes:exportar"),
            "linhas": [
                linha_da_lista(s, permissions.acoes_permitidas(request.user, s))
                for s in pagina
            ],
        },
    )


COLUNAS_EXPORTACAO = [
    ("Nº", 8), ("Status", 22), ("Data da solicitação", 14), ("Início do evento", 14),
    ("Fim do evento", 14), ("Município", 22), ("Região", 18), ("Tipo de evento", 24),
    ("Local", 30), ("Protocolo", 14), ("Solicitante", 28), ("Cargo / unidade", 28),
    ("Contato", 16), ("Órgão responsável", 24), ("Serviços", 36),
    ("Equipes (servidores)", 36), ("Total de servidores", 12), ("Tipo de operação", 14),
    ("Unidade móvel", 10), ("Qtde CIN", 10), ("Motorista", 22), ("Decisão DG", 16),
    ("Observações DG", 40), ("Decidido por", 20), ("Decidido em", 17), ("Criado por", 20),
]
# Colunas de data (índice a partir de 0) e a data/hora da decisão.
COLUNAS_DATA = {2, 3, 4}
COLUNA_DATA_HORA = 24


def _linhas_exportacao(queryset):
    """Uma linha por solicitação, nos tipos de verdade (data é data, número é número)."""
    from django.utils import timezone as tz

    for s in queryset:
        equipes = "; ".join(
            f"{item.equipe} ({item.quantidade_servidores or '-'})"
            for item in s.itens_equipe.all()
        )
        yield [
            s.pk,
            s.get_status_display(),
            s.data_solicitacao,
            s.data_inicio_evento,
            s.data_fim_evento,
            str(s.municipio or ""),
            str(s.regiao or ""),
            str(s.tipo_evento or ""),
            s.local_evento,
            s.protocolo,
            s.solicitante_nome,
            s.solicitante_cargo_unidade,
            s.contato,
            str(s.orgao_responsavel or ""),
            "; ".join(str(servico) for servico in s.servicos.all()),
            equipes,
            s.quantidade_servidores,
            s.get_tipo_operacao_display() if s.tipo_operacao else "",
            "Sim" if s.unidade_movel else "Não",
            s.quantidade_cin,
            str(s.motorista or ""),
            s.get_decisao_dg_display(),
            s.observacoes_dg,
            str(s.decidido_por or ""),
            tz.localtime(s.decidido_em).replace(tzinfo=None) if s.decidido_em else None,
            str(s.criado_por),
        ]


def _exportar_csv(linhas, nome):
    import csv

    from django.http import HttpResponse

    resposta = HttpResponse(content_type="text/csv; charset=utf-8")
    resposta["Content-Disposition"] = f'attachment; filename="{nome}.csv"'
    # BOM para o Excel reconhecer UTF-8; ponto e vírgula para o Excel pt-BR.
    resposta.write("\ufeff")
    escritor = csv.writer(resposta, delimiter=";", lineterminator="\r\n")
    escritor.writerow([titulo for titulo, _largura in COLUNAS_EXPORTACAO])
    for linha in linhas:
        escritor.writerow([
            ""
            if valor is None
            else valor.strftime("%d/%m/%Y %H:%M" if indice == COLUNA_DATA_HORA else "%d/%m/%Y")
            if hasattr(valor, "strftime")
            else valor
            for indice, valor in enumerate(linha)
        ])
    return resposta


def _exportar_xlsx(linhas, nome):
    """Planilha formatada: cabeçalho fixo com filtro, datas como datas."""
    from io import BytesIO

    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    aba = livro.active
    aba.title = "Solicitações"
    aba.append([titulo for titulo, _largura in COLUNAS_EXPORTACAO])
    for celula in aba[1]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="D1D3D4")
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    for linha in linhas:
        aba.append(linha)
        # Texto que começa com "=" viraria fórmula: fica texto.
        for celula in aba[aba.max_row]:
            if celula.data_type == "f":
                celula.data_type = "s"
    for indice, (_titulo, largura) in enumerate(COLUNAS_EXPORTACAO):
        letra = get_column_letter(indice + 1)
        aba.column_dimensions[letra].width = largura
        if indice in COLUNAS_DATA or indice == COLUNA_DATA_HORA:
            formato = "DD/MM/YYYY HH:MM" if indice == COLUNA_DATA_HORA else "DD/MM/YYYY"
            for (celula,) in aba.iter_rows(min_row=2, min_col=indice + 1, max_col=indice + 1):
                celula.number_format = formato
    aba.freeze_panes = "B2"
    aba.auto_filter.ref = aba.dimensions
    saida = BytesIO()
    livro.save(saida)
    resposta = HttpResponse(
        saida.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resposta["Content-Disposition"] = f'attachment; filename="{nome}.xlsx"'
    return resposta


@login_required
def exportar_solicitacoes(request):
    """Exporta a lista com a fila, a busca e os filtros da tela.

    Sai em .xlsx formatado; `?formato=csv` mantém o CSV de antes (links e
    rotinas que já o usavam).
    """
    queryset, _base, _filtros, _fila = _queryset_filtrado(request)
    queryset = queryset.select_related(
        "orgao_responsavel", "motorista", "decidido_por"
    ).prefetch_related("servicos", "itens_equipe__equipe")
    nome = f"solicitacoes-{timezone.localdate():%Y-%m-%d}"
    linhas = _linhas_exportacao(queryset)
    if request.GET.get("formato") == "csv":
        return _exportar_csv(linhas, nome)
    return _exportar_xlsx(linhas, nome)


DECISOES_DG = [
    ("ATENDER", "Atender", "Deferida — em andamento; o solicitante confirma depois do evento", "check-circle"),
    ("NAO_ATENDER", "Não atender", "Encerra como não atendida; observação obrigatória", "x"),
    ("CANCELADO", "Evento cancelado", "Encerra como cancelada; observação obrigatória", "ban"),
    ("DEVOLVER", "Enviar para correção", "Volta ao solicitante; diga brevemente o que ele deve mudar", "undo"),
]


def _decisoes_dg(pendente):
    """As decisões da DG no contrato dos cartões de escolha."""
    escolhida = (pendente or {}).get("decisao", "")
    return [
        {"valor": valor, "rotulo": rotulo, "dica": dica, "icone": icone,
         "marcado": escolhida == valor}
        for valor, rotulo, dica, icone in DECISOES_DG
    ]


def _despacho_pendente(request, solicitacao):
    """Decisão e observação que o gestor tentou registrar e não passaram."""
    pendente = request.session.pop("despacho_pendente", None)
    if not pendente or pendente.get("solicitacao") != solicitacao.pk:
        return None
    return pendente


# ---------------------------------------------------------------------------
# Transições de workflow (somente POST)
# ---------------------------------------------------------------------------

def _voltar_ao_despacho(request, solicitacao, decisao="", observacao=""):
    """Devolve o gestor à seção do despacho com o que ele já tinha escolhido."""
    request.session["despacho_pendente"] = {
        "solicitacao": solicitacao.pk,
        "decisao": decisao,
        "observacao": observacao,
    }
    url = reverse("solicitacoes:editar", args=[solicitacao.pk])
    return redirect(f"{url}#despacho-dg")


def _executar_transicao(request, solicitacao, funcao, mensagem, **kwargs):
    try:
        funcao(solicitacao, request.user, **kwargs)
    except ValidationError as erro:
        for mensagem_erro in erro.messages:
            messages.error(request, mensagem_erro)
    else:
        messages.success(request, mensagem)
    return redirect("solicitacoes:editar", pk=solicitacao.pk)


@login_required
@require_POST
def enviar_solicitacao(request, pk):
    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_enviar(request.user, solicitacao):
        raise PermissionDenied
    return _executar_transicao(
        request, solicitacao, services.enviar,
        f"Solicitação #{solicitacao.pk} enviada com sucesso.",
    )


@login_required
@require_POST
def concluir_solicitacao(request, pk):
    """O solicitante confirma que o evento aconteceu e foi atendido."""
    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_concluir(request.user, solicitacao):
        raise PermissionDenied
    return _executar_transicao(
        request, solicitacao, services.concluir_atendimento,
        f"Atendimento da solicitação #{solicitacao.pk} confirmado.",
    )


@login_required
@require_POST
def cancelar_evento(request, pk):
    """Criador ou perfil institucional autorizado registra o cancelamento."""
    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_cancelar(request.user, solicitacao):
        raise PermissionDenied
    return _executar_transicao(
        request, solicitacao, services.cancelar_evento,
        f"Evento da solicitação #{solicitacao.pk} registrado como cancelado.",
        observacao=request.POST.get("motivo_cancelamento", ""),
    )


def _opcoes_responsaveis(solicitacao):
    """Usuários ativos que podem assumir a solicitação (menos o atual)."""
    from django.contrib.auth import get_user_model

    usuarios = (
        get_user_model()
        .objects.filter(is_active=True)
        .exclude(pk=solicitacao.criado_por_id)
        .order_by("first_name", "last_name", "username")
    )
    return [{"valor": str(u.pk), "rotulo": str(u)} for u in usuarios]


@login_required
@require_POST
def transferir_solicitacao(request, pk):
    """Passa a solicitação para outro responsável (histórico e aviso no sino)."""
    from django.contrib.auth import get_user_model

    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_transferir(request.user, solicitacao):
        raise PermissionDenied
    escolhido = str(request.POST.get("responsavel", "")).strip()
    novo = (
        get_user_model().objects.filter(pk=escolhido).first()
        if escolhido.isdigit()
        else None
    )
    try:
        services.transferir(
            solicitacao, request.user, novo, request.POST.get("motivo_transferencia", "")
        )
    except ValidationError as erro:
        for mensagem_erro in erro.messages:
            messages.error(request, mensagem_erro)
        url = reverse("solicitacoes:editar", args=[solicitacao.pk])
        return redirect(f"{url}#responsavel")
    messages.success(
        request, f"Solicitação #{solicitacao.pk} transferida para {novo}."
    )
    # Quem passou adiante pode deixar de enxergar o pedido.
    if not permissions.pode_ver(request.user, solicitacao):
        return redirect("solicitacoes:lista")
    return redirect("solicitacoes:editar", pk=solicitacao.pk)


@login_required
@require_POST
def consultar_protocolo(request, pk):
    """"Consultar andamento": a última movimentação do protocolo no eProtocolo.

    Só leitura; o resultado volta num cartão da própria tela.
    """
    solicitacao = _obter_visivel(request, pk)
    if not solicitacao.protocolo:
        messages.error(request, "Informe e salve o número do protocolo antes de consultar.")
    else:
        erro = andamento_eprotocolo.consultar_e_guardar(
            request, "solicitacoes", solicitacao.pk, solicitacao.protocolo
        )
        if erro:
            messages.error(request, erro)
    url = reverse("solicitacoes:editar", args=[solicitacao.pk])
    return redirect(f"{url}#protocolo")


@login_required
@require_POST
def gerar_viagem(request, pk):
    """"Gerar viagem": para a deferida que ficou sem viagem.

    Vale para as deferidas antes da geração automática e para quando ela
    falhou no despacho. Liberado à DG e a quem opera Viagens — este último
    pode não enxergar a solicitação, e então vai direto para a viagem criada.
    """
    solicitacao = get_object_or_404(SolicitacaoEvento, pk=pk)
    if not permissions.pode_gerar_viagem(request.user):
        raise PermissionDenied
    ve_a_solicitacao = permissions.pode_ver(request.user, solicitacao)
    try:
        viagem = integracao_viagens.gerar_viagem(solicitacao, request.user)
    except ValueError as erro:
        messages.error(request, f"Não foi possível gerar a viagem: {erro}")
        if not ve_a_solicitacao:
            return redirect("viagens_viagem:lista")
    else:
        messages.success(
            request,
            f"Viagem #{viagem.pk} gerada em rascunho. Complete sede, horário, "
            "servidores, viatura e motorista em Viagens.",
        )
        if not ve_a_solicitacao:
            return redirect("viagens_viagem:painel", pk=viagem.pk)
    url = reverse("solicitacoes:editar", args=[solicitacao.pk])
    return redirect(f"{url}#viagem")


# ---------------------------------------------------------------------------
# Anexos
# ---------------------------------------------------------------------------

@login_required
def adicionar_anexo(request, pk):
    """Anexa um arquivo — pelo modal da tela ou por POST direto.

    No modal vale o protocolo dos cadastros (`X-Cadastro-Modal`): o GET
    devolve o trecho do formulário e o POST devolve `{"ok": true}` ou o
    trecho de novo, com o erro.
    """
    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_gerenciar_anexos(request.user, solicitacao):
        raise PermissionDenied
    via_modal = request.headers.get("X-Cadastro-Modal") == "1"
    if request.method != "POST":
        if not via_modal:
            return redirect("solicitacoes:editar", pk=solicitacao.pk)
        return render(
            request,
            "pages/solicitacoes/_modal_anexo.html",
            {"solicitacao": solicitacao},
        )
    form = AnexoForm(request.POST, request.FILES)
    if form.is_valid():
        arquivo = form.cleaned_data["arquivo"]
        AnexoSolicitacao.objects.create(
            solicitacao=solicitacao,
            arquivo=arquivo,
            nome_original=arquivo.name,
            tamanho=arquivo.size,
            enviado_por=request.user,
        )
        services.registrar_historico(
            solicitacao, request.user, AcaoHistorico.ATUALIZACAO,
            status_novo=solicitacao.status,
            observacao=f"Anexo adicionado: {arquivo.name}",
        )
        messages.success(request, f"Arquivo {arquivo.name} anexado.")
        if via_modal:
            return JsonResponse({"ok": True})
    else:
        mensagens = [erro for erros in form.errors.values() for erro in erros]
        if via_modal:
            return render(
                request,
                "pages/solicitacoes/_modal_anexo.html",
                {"solicitacao": solicitacao, "erro_anexo": " ".join(mensagens)},
            )
        for mensagem in mensagens:
            messages.error(request, mensagem)
    return redirect("solicitacoes:editar", pk=solicitacao.pk)


@login_required
def baixar_anexo(request, pk, anexo_pk):
    solicitacao = _obter_visivel(request, pk)
    anexo = get_object_or_404(solicitacao.anexos, pk=anexo_pk)
    return FileResponse(
        anexo.arquivo.open("rb"),
        as_attachment=True,
        filename=anexo.nome_original,
    )


@login_required
@require_POST
def excluir_anexo(request, pk, anexo_pk):
    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_gerenciar_anexos(request.user, solicitacao):
        raise PermissionDenied
    anexo = get_object_or_404(solicitacao.anexos, pk=anexo_pk)
    nome = anexo.nome_original
    anexo.delete()
    services.registrar_historico(
        solicitacao, request.user, AcaoHistorico.ATUALIZACAO,
        status_novo=solicitacao.status,
        observacao=f"Anexo removido: {nome}",
    )
    messages.success(request, f"Anexo {nome} removido.")
    return redirect("solicitacoes:editar", pk=solicitacao.pk)


@login_required
@require_POST
def excluir_solicitacao(request, pk):
    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_excluir(request.user, solicitacao):
        raise PermissionDenied
    numero = solicitacao.pk
    solicitacao.delete()
    messages.success(request, f"Rascunho #{numero} excluído.")
    return redirect("solicitacoes:lista")


def _quantidades_dg_do_post(request, solicitacao):
    """Quantidades por equipe informadas pela DG no formulário de despacho."""
    quantidades = {}
    for item in solicitacao.itens_equipe.all():
        valor = str(request.POST.get(f"quantidade_dg_{item.equipe_id}", "")).strip()
        if valor:
            try:
                quantidades[item.equipe_id] = int(valor)
            except ValueError:
                quantidades[item.equipe_id] = 0
    return quantidades


@login_required
@require_POST
def despachar(request, pk):
    solicitacao = _obter_visivel(request, pk)
    if not permissions.pode_despachar(request.user, solicitacao):
        raise PermissionDenied

    # Salvar apenas os ajustes de quantidade, sem decidir ainda.
    if request.POST.get("acao_despacho") == "salvar_ajustes":
        try:
            mudancas = services.salvar_ajustes_dg(
                solicitacao, request.user, _quantidades_dg_do_post(request, solicitacao)
            )
        except ValidationError as erro:
            for mensagem_erro in erro.messages:
                messages.error(request, mensagem_erro)
        else:
            if mudancas:
                messages.success(
                    request,
                    f"Ajustes salvos para a solicitação #{solicitacao.pk}: "
                    + "; ".join(mudancas) + ".",
                )
            else:
                messages.info(request, "Nenhuma alteração nas quantidades.")
        return redirect("solicitacoes:editar", pk=solicitacao.pk)

    form = DespachoForm(request.POST)
    if not form.is_valid():
        for erros_campo in form.errors.values():
            for erro in erros_campo:
                messages.error(request, erro)
        return _voltar_ao_despacho(
            request,
            solicitacao,
            request.POST.get("decisao", ""),
            request.POST.get("observacao", ""),
        )

    decisao = form.cleaned_data["decisao"]
    observacao = form.cleaned_data["observacao"]
    # "Registrar e abrir a próxima": a ordem é a de antes da decisão.
    seguir = request.POST.get("seguir") == "proxima"
    fila = _fila_de_despacho(request.user) if seguir else []
    try:
        if decisao == DespachoForm.DEVOLVER:
            services.devolver(solicitacao, request.user, observacao=observacao)
            sucesso = f"Solicitação #{solicitacao.pk} enviada para correção."
        else:
            # A DG aceita as quantidades propostas ou informa novas por equipe.
            services.despachar(
                solicitacao,
                request.user,
                decisao=decisao,
                observacao=observacao,
                quantidades=_quantidades_dg_do_post(request, solicitacao),
            )
            sucesso = f"Decisão registrada para a solicitação #{solicitacao.pk}."
    except ValidationError as erro:
        for mensagem_erro in erro.messages:
            messages.error(request, mensagem_erro)
        return _voltar_ao_despacho(request, solicitacao, decisao, observacao)
    except DatabaseError:
        # A transação do serviço já foi desfeita: nada ficou pela metade. A
        # DG volta ao despacho com o que escreveu, em vez de uma página 500.
        logger.exception("Falha ao gravar o despacho da solicitação %s.", solicitacao.pk)
        messages.error(
            request,
            "Não foi possível gravar o despacho agora. Nada foi alterado; "
            "confira o texto e tente de novo.",
        )
        return _voltar_ao_despacho(request, solicitacao, decisao, observacao)

    messages.success(request, sucesso)
    if seguir:
        return _abrir_proxima_do_despacho(request, solicitacao.pk, fila)
    return redirect("solicitacoes:editar", pk=solicitacao.pk)


def _abrir_proxima_do_despacho(request, atual, fila):
    """Depois da decisão, a próxima da fila (a que vinha depois da atual)."""
    depois = fila[fila.index(atual) + 1:] if atual in fila else fila
    antes = fila[:fila.index(atual)] if atual in fila else []
    ainda_pendentes = set(
        SolicitacaoEvento.objects.filter(
            pk__in=fila, status=StatusSolicitacao.AGUARDANDO_DESPACHO
        ).values_list("pk", flat=True)
    )
    proxima = next((pk for pk in depois + antes if pk in ainda_pendentes), None)
    if proxima is None:
        messages.info(request, "Não há mais solicitações aguardando despacho.")
        return redirect(f"{reverse('solicitacoes:lista')}?fila=despacho")
    url = reverse("solicitacoes:editar", args=[proxima])
    return redirect(f"{url}#despacho-dg")
