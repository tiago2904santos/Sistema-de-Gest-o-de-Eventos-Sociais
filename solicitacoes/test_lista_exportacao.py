"""Lista de solicitações: botão Exportar, painel de filtros e busca ampliada (m006)."""

from datetime import date
from io import BytesIO

from django.urls import reverse

from cadastros.models import OrgaoResponsavel, TipoEvento

from . import services
from .tests import BaseSolicitacaoTestCase


class ExportarPlanilhaTests(BaseSolicitacaoTestCase):
    def _planilha(self, resposta):
        from openpyxl import load_workbook

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(
            resposta["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", resposta["Content-Disposition"])
        return load_workbook(BytesIO(resposta.content)).active

    def test_exporta_xlsx_formatado_com_datas_de_verdade(self):
        solicitacao = self.solicitacao_completa()
        services.enviar(solicitacao, self.solicitante)
        self.client.force_login(self.solicitante)

        aba = self._planilha(self.client.get(reverse("solicitacoes:exportar")))

        self.assertEqual(aba["A1"].value, "Nº")
        self.assertTrue(aba["A1"].font.bold)
        self.assertEqual(aba.freeze_panes, "B2")
        self.assertEqual(aba["A2"].value, solicitacao.pk)
        self.assertEqual(aba["B2"].value, "Aguardando despacho")
        self.assertEqual(aba["D2"].value.date(), date(2026, 9, 10))
        self.assertEqual(aba["D2"].number_format, "DD/MM/YYYY")
        self.assertIn("Equipe Alfa (5)", [c.value for c in aba[2]])

    def test_texto_com_igual_nao_vira_formula(self):
        self.criar_solicitacao(solicitante_nome="=HYPERLINK(\"x\")")
        self.client.force_login(self.solicitante)
        aba = self._planilha(self.client.get(reverse("solicitacoes:exportar")))
        celula = aba["K2"]
        self.assertEqual(celula.data_type, "s")

    def test_exportacao_respeita_a_fila_e_a_busca(self):
        enviada = self.solicitacao_completa()
        services.enviar(enviada, self.solicitante)
        self.criar_solicitacao(local_evento="Escola municipal")
        self.client.force_login(self.solicitante)

        aba = self._planilha(
            self.client.get(reverse("solicitacoes:exportar"), {"fila": "rascunhos"})
        )
        self.assertEqual(aba.max_row, 2)
        self.assertEqual(aba["I2"].value, "Escola municipal")

    def test_botao_exportar_leva_a_querystring_da_tela(self):
        self.client.force_login(self.solicitante)
        resposta = self.client.get(
            reverse("solicitacoes:lista"), {"fila": "minhas", "q": "praça"}
        )
        url = reverse("solicitacoes:exportar")
        self.assertContains(resposta, f'href="{url}?fila=minhas&amp;q=pra%C3%A7a"')

    def test_lista_de_palestras_tambem_tem_o_botao(self):
        self.client.force_login(self.superusuario)
        resposta = self.client.get(reverse("demandas_eventos:lista"))
        self.assertContains(resposta, reverse("demandas_eventos:exportar"))


class PainelFiltrosTests(BaseSolicitacaoTestCase):
    def test_painel_recolhido_sem_filtro(self):
        self.criar_solicitacao()
        self.client.force_login(self.solicitante)
        resposta = self.client.get(reverse("solicitacoes:lista"))
        self.assertContains(resposta, 'class="filtros-painel"')
        self.assertNotContains(resposta, 'class="filtros-painel" open')
        self.assertContains(resposta, 'name="tipo_evento"')

    def test_painel_abre_com_filtro_e_guarda_a_fila(self):
        self.criar_solicitacao()
        self.client.force_login(self.solicitante)
        resposta = self.client.get(
            reverse("solicitacoes:lista"),
            {"fila": "minhas", "municipio": self.municipio.pk, "inicio": "2026-09-01"},
        )
        painel = resposta.context["painel"]
        self.assertTrue(painel["aberto"])
        self.assertEqual(painel["total"], 2)
        self.assertIn({"nome": "fila", "valor": "minhas"}, painel["ocultos"])
        self.assertEqual(painel["url_limpar"], "?fila=minhas")
        self.assertEqual(len(resposta.context["linhas"]), 1)

    def test_filtro_por_tipo_e_periodo(self):
        outro_tipo = TipoEvento.objects.create(nome="Feira de saúde")
        self.criar_solicitacao(tipo_evento=outro_tipo)
        self.criar_solicitacao(data_inicio_evento=date(2026, 12, 1), data_fim_evento=date(2026, 12, 1))
        self.client.force_login(self.solicitante)
        resposta = self.client.get(reverse("solicitacoes:lista"), {"tipo_evento": outro_tipo.pk})
        self.assertEqual(len(resposta.context["linhas"]), 1)
        resposta = self.client.get(reverse("solicitacoes:lista"), {"inicio": "2026-11-01"})
        self.assertEqual(len(resposta.context["linhas"]), 1)


class BuscaAmpliadaTests(BaseSolicitacaoTestCase):
    def _achados(self, termo):
        self.client.force_login(self.solicitante)
        resposta = self.client.get(reverse("solicitacoes:lista"), {"q": termo})
        return [linha["solicitacao"].pk for linha in resposta.context["linhas"]]

    def test_busca_pelo_numero(self):
        alvo = self.criar_solicitacao()
        self.criar_solicitacao()
        self.assertEqual(self._achados(f"#{alvo.pk}"), [alvo.pk])
        self.assertIn(alvo.pk, self._achados(str(alvo.pk)))

    def test_busca_pelo_tipo_e_pelo_orgao(self):
        tipo = TipoEvento.objects.create(nome="Feira de saúde")
        orgao = OrgaoResponsavel.objects.create(nome="Prefeitura de Exemplo")
        por_tipo = self.criar_solicitacao(tipo_evento=tipo)
        por_orgao = self.criar_solicitacao(orgao_responsavel=orgao)
        self.assertEqual(self._achados("feira"), [por_tipo.pk])
        self.assertEqual(self._achados("prefeitura de ex"), [por_orgao.pk])
