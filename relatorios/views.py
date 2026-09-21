"""O relatório consolidado: a tela e a planilha, com os mesmos números."""

from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from .consolidacao import anos_disponiveis, relatorio


def _ano(request, hoje):
    """O ano pedido; "todos" junta tudo; qualquer outra coisa vira o ano atual."""
    valor = request.GET.get("ano", "").strip()
    if valor == "todos":
        return None
    if valor.isdigit() and int(valor) in anos_disponiveis(hoje):
        return int(valor)
    return hoje.year


@login_required
def painel(request):
    hoje = timezone.localdate()
    ano = _ano(request, hoje)
    dados = relatorio(request.user, ano, hoje)
    return render(
        request,
        "pages/relatorios/painel.html",
        {
            **dados,
            "ano_valor": str(ano) if ano else "todos",
            "opcoes_ano": [{"valor": str(a), "rotulo": str(a)} for a in anos_disponiveis(hoje)]
            + [{"valor": "todos", "rotulo": "Todos os anos"}],
            "gerado_em": timezone.localtime(),
        },
    )


@login_required
def exportar(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    hoje = timezone.localdate()
    ano = _ano(request, hoje)
    dados = relatorio(request.user, ano, hoje)

    livro = Workbook()
    livro.remove(livro.active)
    negrito = Font(bold=True)
    cabecalho = PatternFill("solid", fgColor="D1D3D4")
    titulo_fonte = Font(bold=True, size=13)

    for secao in [dados["quadro"], *dados["secoes"]]:
        aba = livro.create_sheet(secao["titulo"][:31])
        aba.append([secao["titulo"].upper() + f" — {dados['periodo'].rotulo}"])
        aba["A1"].font = titulo_fonte
        aba.append(secao["colunas"])
        for celula in aba[2]:
            celula.font = negrito
            celula.fill = cabecalho
        for linha in secao["linhas"]:
            aba.append(linha)
        if secao["totais"]:
            aba.append(secao["totais"])
            for celula in aba[aba.max_row]:
                celula.font = negrito
        if secao["nota"]:
            aba.append([])
            aba.append([secao["nota"]])
            aba.cell(aba.max_row, 1).alignment = Alignment(wrap_text=False)
            aba.cell(aba.max_row, 1).font = Font(italic=True, color="808080")
        for indice, coluna in enumerate(secao["colunas"], start=1):
            maior = max(
                [len(str(coluna))]
                + [len(str(linha[indice - 1])) for linha in secao["linhas"] if indice - 1 < len(linha)]
            )
            aba.column_dimensions[get_column_letter(indice)].width = min(max(maior + 2, 9), 60)
        aba.freeze_panes = "A3"

    saida = BytesIO()
    livro.save(saida)
    nome = f"relatorio-consolidado-{ano or 'todos-os-anos'}.xlsx"
    resposta = HttpResponse(
        saida.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resposta["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resposta
