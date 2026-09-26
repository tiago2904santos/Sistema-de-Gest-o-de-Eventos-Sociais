"""O relatório consolidado — quem vê o quê, e se os números fecham."""

from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Setor
from cadastros.models import TipoEvento
from demandas_eventos.models import DemandaEvento, StatusDemanda, TipoEventoPalestra
from solicitacoes.models import StatusSolicitacao
from solicitacoes.tests import BaseSolicitacaoTestCase

from .consolidacao import relatorio

User = get_user_model()
HOJE = date(2026, 9, 21)


class RelatorioConsolidadoTests(BaseSolicitacaoTestCase):
    def setUp(self):
        super().setUp()
        self.ascom = Setor.objects.get(sigla="ASCOM")
        self.da_ascom = User.objects.create_user("ascom_rel", password="x")
        self.da_ascom.setores.add(self.ascom)
        self.pcpr = TipoEvento.objects.get_or_create(nome="PCPR na Comunidade")[0]

    def palestra(self, **kwargs):
        dados = {
            "data_solicitacao": date(2026, 3, 1),
            "data_inicio_evento": date(2026, 3, 10),
            "evento": TipoEventoPalestra.PALESTRA,
            "status": StatusDemanda.ATENDIDA,
            "solicitante": "Escola",
            "quantidade_publico": 100,
        }
        dados.update(kwargs)
        d = DemandaEvento.objects.create(**dados)
        d.setores.add(self.ascom)
        return d

    def test_anonimo_vai_para_o_login(self):
        resposta = self.client.get(reverse("relatorios:painel"))
        self.assertEqual(resposta.status_code, 302)

    def test_palestras_por_mes_contam_so_as_atendidas(self):
        self.palestra()
        self.palestra(quantidade_publico=50)
        self.palestra(status=StatusDemanda.CANCELADA)
        self.palestra(data_inicio_evento=date(2025, 3, 10))
        dados = relatorio(self.da_ascom, 2026, HOJE)
        palestras = next(s for s in dados["secoes"] if s["slug"] == "palestras")
        marco = palestras["linhas"][2]
        self.assertEqual(marco, ["mar/26", 2, 150])
        self.assertEqual(palestras["totais"], ["Total", 2, 150])

    def test_secao_de_modulo_sem_acesso_nao_aparece(self):
        self.palestra()
        dados = relatorio(self.solicitante, 2026, HOJE)
        slugs = {s["slug"] for s in dados["secoes"]}
        self.assertNotIn("palestras", slugs)
        self.assertIn("pcpr", slugs)
        self.client.force_login(self.solicitante)
        resposta = self.client.get(reverse("relatorios:painel"))
        self.assertEqual(resposta.status_code, 200)
        self.assertNotContains(resposta, 'id="rel-palestras"')

    def test_pcpr_junta_as_duas_fontes_sem_duplicar(self):
        solicitacao = self.criar_solicitacao(criado_por=self.da_ascom)
        solicitacao.tipo_evento = self.pcpr
        solicitacao.status = StatusSolicitacao.ATENDIDA
        solicitacao.data_inicio_evento = date(2026, 5, 9)
        solicitacao.data_fim_evento = date(2026, 5, 10)
        solicitacao.quantidade_cin = 230
        solicitacao.save()
        self.palestra(
            evento=TipoEventoPalestra.PCPR_NA_COMUNIDADE,
            municipio=solicitacao.municipio,
            data_inicio_evento=date(2026, 5, 10),
            quantidade_publico=900,
        )
        dados = relatorio(self.da_ascom, 2026, HOJE)
        pcpr = next(s for s in dados["secoes"] if s["slug"] == "pcpr")
        self.assertEqual(len(pcpr["linhas"]), 1)
        linha = pcpr["linhas"][0]
        self.assertEqual(linha[3:], [900, 230, "Solicitação + ASCOM"])

    def test_pcpr_casa_primeiro_pelo_vinculo_do_encaminhamento(self):
        solicitacao = self.criar_solicitacao(criado_por=self.da_ascom)
        solicitacao.tipo_evento = self.pcpr
        solicitacao.status = StatusSolicitacao.ATENDIDA
        solicitacao.data_inicio_evento = date(2026, 5, 9)
        solicitacao.data_fim_evento = date(2026, 5, 9)
        solicitacao.quantidade_cin = 120
        solicitacao.save()
        # Sem município e com a data remarcada na planilha: só o vínculo junta as duas.
        self.palestra(
            evento=TipoEventoPalestra.PCPR_NA_COMUNIDADE,
            data_inicio_evento=date(2026, 5, 16),
            quantidade_publico=400,
            solicitacao_dg=solicitacao,
        )
        dados = relatorio(self.da_ascom, 2026, HOJE)
        pcpr = next(s for s in dados["secoes"] if s["slug"] == "pcpr")
        self.assertEqual(len(pcpr["linhas"]), 1)
        self.assertEqual(pcpr["linhas"][0][3:], [400, 120, "Solicitação + ASCOM"])

    def test_painel_e_planilha_trazem_as_mesmas_secoes(self):
        self.palestra()
        self.client.force_login(self.da_ascom)
        resposta = self.client.get(reverse("relatorios:painel"), {"ano": "2026"})
        self.assertContains(resposta, "Relatório consolidado")
        self.assertContains(resposta, 'id="rel-palestras"')
        resposta = self.client.get(reverse("relatorios:exportar"), {"ano": "2026"})
        self.assertEqual(resposta.status_code, 200)
        livro = load_workbook(BytesIO(resposta.content))
        self.assertIn("Palestras", livro.sheetnames)
        self.assertIn("PCPR na Comunidade", livro.sheetnames)
        self.assertIn("Visão geral por mês", livro.sheetnames)

    def test_todos_os_anos_e_ano_invalido(self):
        self.palestra(data_inicio_evento=date(2024, 6, 1))
        self.client.force_login(self.da_ascom)
        self.assertEqual(self.client.get(reverse("relatorios:painel"), {"ano": "todos"}).status_code, 200)
        self.assertEqual(self.client.get(reverse("relatorios:painel"), {"ano": "abc"}).status_code, 200)
        dados = relatorio(self.da_ascom, None, HOJE)
        palestras = next(s for s in dados["secoes"] if s["slug"] == "palestras")
        self.assertIn(["2024", 1, 100], palestras["linhas"])
